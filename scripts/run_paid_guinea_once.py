"""Run exactly one paid Guinea assessment without assistant follow-up calls."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from time import monotonic, sleep
from urllib.parse import urlsplit
from zipfile import ZipFile

from docx import Document
from playwright.sync_api import sync_playwright

SITE = "https://cpf-fcv-review-prototype.onrender.com/"
FCV_PATTERN = re.compile(
    r"\\b(?:conflicts?|violence|violent|political|governance|government|elections?|coup|"
    r"humanitarian|displacement|displaced|refugees?|protection|peace|peacebuilding|"
    r"insecurity)\\b|\\bland (?:conflict|dispute|tenure)\\b|\\bsocial cohesion\\b|"
    r"\\bsecurity (?:update|situation|incident|threat|forces?|sector|crisis|risk)\\b",
    re.IGNORECASE,
)
THEMES = {
    "political": ("political", "transition", "election", "coup", "governance", "government"),
    "conflict": ("conflict", "violence", "violent", "security", "insecurity", "peace"),
    "displacement": ("displacement", "displaced", "refugee", "humanitarian", "protection"),
    "land": ("land conflict", "land dispute", "land tenure"),
}


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, indent=2, sort_keys=True), encoding="utf-8")


def themes(text: str) -> set[str]:
    normalized = text.casefold()
    return {
        name for name, markers in THEMES.items() if any(marker in normalized for marker in markers)
    }


def clean_identifier_keys(value: object) -> object:
    if isinstance(value, dict):
        return {
            key: clean_identifier_keys(item)
            for key, item in value.items()
            if key != "assessment_id"
        }
    if isinstance(value, list):
        return [clean_identifier_keys(item) for item in value]
    return value


def validate_docx(path: Path) -> dict[str, int]:
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= names
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for heading in (
        "Overall assessment",
        "RRA driver-to-response assessment",
        "Priority areas for strengthening",
        "Basis and important limitations",
    ):
        assert heading in text
    return {"bytes": path.stat().st_size, "paragraphs": len(document.paragraphs)}


def summarize(payload: dict[str, object], docx: Path) -> dict[str, object]:
    evidence = payload["evidence_by_id"]
    result = {key: value for key, value in payload.items() if key != "evidence_by_id"}
    current = {
        key: item for key, item in evidence.items() if item.get("evidence_type") == "current_context"
    }
    current_hosts = Counter(
        urlsplit(item.get("source_url") or "").hostname or "missing" for item in current.values()
    )
    indicator_count = sum(host == "api.worldbank.org" for host in current_hosts.elements())
    relevant_current = {
        key for key, item in current.items() if FCV_PATTERN.search(item.get("text") or "")
    }
    mismatched_current_citations = 0
    current_citation_count = 0
    pathways_present = 0
    priorities = result.get("priority_areas") or []
    for priority in priorities:
        priority_text = " ".join(
            str(priority.get(field) or "")
            for field in ("heading", "assessment", "why_it_matters", "recommended_action")
        )
        priority_themes = themes(priority_text)
        why = str(priority.get("why_it_matters") or "").casefold()
        if any(marker in why for marker in ("because", "through", "thereby", "which", "undermin", "exacerbat", "reduce", "mitigate", "prevent", "strengthen")):
            pathways_present += 1
        for evidence_id in priority.get("evidence_ids") or []:
            if evidence_id not in current:
                continue
            current_citation_count += 1
            evidence_themes = themes(str(current[evidence_id].get("text") or ""))
            if not evidence_themes or not (evidence_themes & priority_themes):
                mismatched_current_citations += 1
    metadata = result.get("metadata") or {}
    return {
        "release": metadata.get("app_release"),
        "tier": metadata.get("current_evidence_tier"),
        "repair_count": metadata.get("repair_count"),
        "validation_outcomes": metadata.get("validation_outcomes"),
        "evidence_count": len(evidence),
        "current_evidence_count": len(current),
        "current_hosts": dict(sorted(current_hosts.items())),
        "fcv_relevant_current_count": len(relevant_current),
        "indicator_api_current_count": indicator_count,
        "priority_count": len(priorities),
        "priority_current_citation_count": current_citation_count,
        "mismatched_current_citation_count": mismatched_current_citations,
        "priorities_with_explicit_pathway": pathways_present,
        "rra_assessment_count": len(result.get("rra_driver_assessments") or []),
        "fcv_strategy_assessment_count": len(result.get("fcv_strategy_assessments") or []),
        "docx": validate_docx(docx),
    }


def main(cpf: Path, rra: Path, output: Path, expected_release: str) -> None:
    output.mkdir(parents=True, exist_ok=False)
    state_path = output / "run-state.json"
    state = {"release": expected_release, "submitted": False, "status": "ready"}
    write_json(state_path, state)
    console_errors: list[str] = []
    page_errors: list[str] = []

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(
            viewport={"width": 1440, "height": 1000}, accept_downloads=True
        )
        page = context.new_page()
        page.on(
            "console",
            lambda message: console_errors.append(message.text)
            if message.type == "error"
            else None,
        )
        page.on("pageerror", lambda error: page_errors.append(str(error)))
        response = page.goto(SITE, wait_until="domcontentloaded", timeout=180_000)
        assert response is not None and response.status == 200
        health = page.request.get(SITE + "health").json()
        assert health["release"] == expected_release, health["release"]

        page.route("**/api/detect-country", lambda route: route.abort())
        page.locator("#cpf").set_input_files(str(cpf))
        confirmation = page.locator(".country-confirmation input")
        confirmation.wait_for(state="visible", timeout=30_000)
        confirmation.fill("Guinea")
        page.locator("#context-documents").set_input_files(str(rra))
        page.locator("#review-stage").select_option("finalization")
        page.wait_for_function(
            "document.querySelector('#country').value === 'Guinea' && "
            "!document.querySelector('#submit-review').disabled",
            timeout=30_000,
        )
        page.screenshot(path=output / "01-intake-full.png", full_page=True)

        state.update({"submitted": True, "status": "submission_started"})
        write_json(state_path, state)
        page.locator("#submit-review").click()
        page.locator("#progress:not([hidden])").wait_for(timeout=10_000)
        page.screenshot(path=output / "02-holding-full.png", full_page=True)

        started = monotonic()
        while monotonic() - started < 1_800:
            if page.locator("#results:not([hidden])").count():
                break
            if page.locator("#research-recovery:not([hidden])").count():
                state["status"] = "research_recovery_required"
                write_json(state_path, state)
                raise RuntimeError("SAFE_FAILURE=research_recovery_required")
            if page.locator("#return-to-intake:not([hidden])").count():
                state["status"] = "run_failed"
                write_json(state_path, state)
                raise RuntimeError("SAFE_FAILURE=run_failed")
            sleep(2)
        else:
            state["status"] = "run_timeout"
            write_json(state_path, state)
            raise TimeoutError("SAFE_FAILURE=run_timeout")

        assessment_id = page.evaluate("sessionStorage.getItem('cpf_fcv_assessment_id')")
        assert assessment_id
        title = page.locator("#result-title").inner_text().strip()
        assert re.fullmatch(r"Guinea (?:CPF|CEN|CPF / CEN) FCV review", title), title
        page.screenshot(path=output / "03-summary-full.png", full_page=True)
        page.locator("#result-tab-detailed").click()
        page.screenshot(path=output / "04-detailed-full.png", full_page=True)

        result_response = page.request.get(SITE + f"api/reviews/{assessment_id}/result")
        assert result_response.status == 200
        payload = clean_identifier_keys(result_response.json())
        write_json(output / "result.json", payload)

        docx = output / "05-full-detailed-note.docx"
        with page.expect_download(timeout=120_000) as download_info:
            page.locator("#export-docx").click()
        download_info.value.save_as(docx)
        summary = summarize(payload, docx)
        assert summary["release"] == expected_release
        assert summary["indicator_api_current_count"] == 0
        assert summary["mismatched_current_citation_count"] == 0
        assert summary["priorities_with_explicit_pathway"] == summary["priority_count"]
        assert not console_errors, console_errors
        assert not page_errors, page_errors
        write_json(output / "quality-summary.json", summary)
        state["status"] = "run_complete"
        write_json(state_path, state)
        print("PAID_RUN_COMPLETE " + json.dumps(summary, sort_keys=True), flush=True)
        context.close()
        browser.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--cpf", type=Path, required=True)
    parser.add_argument("--rra", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--release", required=True)
    args = parser.parse_args()
    main(args.cpf.resolve(), args.rra.resolve(), args.output.resolve(), args.release)
