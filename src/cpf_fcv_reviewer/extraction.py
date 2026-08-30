from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


class DocumentUnreadable(ValueError):
    """Raised when the required primary document contains too little text."""


class ExtractionLimitExceeded(ValueError):
    """Raised when a caller-specific extraction budget is exceeded."""


@dataclass(frozen=True)
class ExtractedSegment:
    text: str
    page: int | None
    heading: str | None
    element: str


@dataclass(frozen=True)
class ExtractedDocument:
    name: str
    segments: tuple[ExtractedSegment, ...]
    warnings: tuple[str, ...]


def segments_from_pdf_pages(name: str, pages: list[str]) -> ExtractedDocument:
    segments: list[ExtractedSegment] = []
    warnings: list[str] = []
    for index, text in enumerate(pages, start=1):
        clean = (text or "").strip()
        if clean:
            segments.append(ExtractedSegment(clean, index, None, f"page {index}"))
        else:
            warnings.append(f"page {index} extracted no text")
    return ExtractedDocument(name, tuple(segments), tuple(warnings))


PDF_SAMPLE_PREFIX_PAGES = 4


def _distributed_page_indices(
    page_count: int, limit: int | None
) -> tuple[int, ...]:
    if limit is not None and limit <= 0:
        return ()
    if limit is None or page_count <= limit:
        return tuple(range(page_count))
    if limit <= 1:
        return (0,)
    prefix_count = min(PDF_SAMPLE_PREFIX_PAGES, limit - 1)
    prefix_indices = tuple(range(prefix_count))
    last_index = page_count - 1
    interior_slots = limit - prefix_count - 1
    if interior_slots <= 0:
        return prefix_indices + (last_index,)
    interior_count = last_index - prefix_count
    if interior_slots == 1:
        interior_indices = (prefix_count,)
    else:
        interior_indices = tuple(
            dict.fromkeys(
                prefix_count
                + round(
                    position * (interior_count - 1) / (interior_slots - 1)
                )
                for position in range(interior_slots)
            )
        )
    return prefix_indices + interior_indices + (last_index,)


def extract_pdf_bytes(
    data: bytes,
    name: str,
    *,
    max_pages: int | None = None,
    sample_across_document: bool = False,
    max_segments: int | None = None,
    max_characters: int | None = None,
    max_uncompressed_bytes: int | None = None,
) -> ExtractedDocument:
    reader = PdfReader(BytesIO(data))
    page_count = len(reader.pages)
    if sample_across_document:
        page_indices = _distributed_page_indices(page_count, max_pages)
    elif max_pages is None or page_count <= max_pages:
        page_indices = tuple(range(page_count))
    else:
        page_indices = tuple(range(max_pages))
    segments: list[ExtractedSegment] = []
    warnings: list[str] = []
    character_count = 0
    uncompressed_bytes = 0

    for page_index in page_indices:
        page = reader.pages[page_index]
        page_number = page_index + 1
        if max_uncompressed_bytes is not None:
            contents = page.get_contents()
            if contents is not None:
                uncompressed_bytes += len(contents.get_data())
                if uncompressed_bytes > max_uncompressed_bytes:
                    raise ExtractionLimitExceeded("PDF stream budget exceeded.")

        clean = (page.extract_text() or "").strip()
        if not clean:
            warnings.append(f"page {page_number} extracted no text")
            continue
        if max_segments is not None and len(segments) >= max_segments:
            raise ExtractionLimitExceeded("PDF segment budget exceeded.")
        next_character_count = character_count + len(clean)
        if max_characters is not None and next_character_count > max_characters:
            raise ExtractionLimitExceeded("PDF character budget exceeded.")
        segments.append(
            ExtractedSegment(
                clean, page_number, None, f"page {page_number}"
            )
        )
        character_count = next_character_count

    if len(page_indices) < page_count:
        warnings.append(
            f"{name}: sampled {len(page_indices)} of {page_count} PDF pages; "
            "conclusions about absence are limited."
        )
    return ExtractedDocument(name, tuple(segments), tuple(warnings))


def extract_docx_bytes(
    data: bytes,
    name: str,
    *,
    max_segments: int | None = None,
    max_characters: int | None = None,
    max_uncompressed_bytes: int | None = None,
    max_archive_members: int | None = None,
) -> ExtractedDocument:
    if max_uncompressed_bytes is not None:
        try:
            with ZipFile(BytesIO(data)) as archive:
                members = archive.infolist()
                if max_archive_members is not None and len(members) > max_archive_members:
                    raise ExtractionLimitExceeded("DOCX archive member budget exceeded.")
                uncompressed_bytes = sum(item.file_size for item in members)
        except (BadZipFile, OSError):
            # Let python-docx produce the normal unreadable-document error below.
            uncompressed_bytes = 0
        if uncompressed_bytes > max_uncompressed_bytes:
            raise ExtractionLimitExceeded("DOCX extraction budget exceeded.")

    document = Document(BytesIO(data))
    segments: list[ExtractedSegment] = []
    warnings: list[str] = []
    heading: str | None = None
    paragraph_number = 0
    table_number = 0
    character_count = 0

    def append_segment(segment: ExtractedSegment) -> None:
        nonlocal character_count
        if max_segments is not None and len(segments) >= max_segments:
            raise ExtractionLimitExceeded("Document segment budget exceeded.")
        next_character_count = character_count + len(segment.text)
        if max_characters is not None and next_character_count > max_characters:
            raise ExtractionLimitExceeded("Document character budget exceeded.")
        segments.append(segment)
        character_count = next_character_count

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            paragraph_number += 1
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                heading = text
            append_segment(
                ExtractedSegment(
                    text=text,
                    page=None,
                    heading=heading,
                    element=f"paragraph {paragraph_number}",
                )
            )
        elif child.tag.endswith("}tbl"):
            table_number += 1
            table = Table(child, document)
            for row_number, row in enumerate(table.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                text = " | ".join(value for value in values if value)
                if text:
                    append_segment(
                        ExtractedSegment(
                            text=text,
                            page=None,
                            heading=heading,
                            element=f"table {table_number} row {row_number}",
                        )
                    )

    if document.inline_shapes:
        warnings.append(
            f"{len(document.inline_shapes)} figure(s) detected; "
            "upload a readable text/table version if they contain material evidence"
        )

    return ExtractedDocument(name, tuple(segments), tuple(warnings))


def require_readable_primary(document: ExtractedDocument) -> None:
    character_count = sum(len(segment.text) for segment in document.segments)
    if character_count < 100:
        raise DocumentUnreadable("Primary CPF/CEN is unreadable or contains too little text.")


TEXT_SEGMENT_CHARACTERS = 4000


def _chunk_text(text: str) -> tuple[str, ...]:
    chunks: list[str] = []
    remaining = text.strip()
    while remaining:
        if len(remaining) <= TEXT_SEGMENT_CHARACTERS:
            chunks.append(remaining)
            break
        window = remaining[: TEXT_SEGMENT_CHARACTERS + 1]
        minimum_boundary = TEXT_SEGMENT_CHARACTERS // 2
        boundary = max(
            window.rfind("\n\n", minimum_boundary),
            window.rfind("\n", minimum_boundary),
            window.rfind(" ", minimum_boundary),
        )
        if boundary < minimum_boundary:
            boundary = TEXT_SEGMENT_CHARACTERS
        chunks.append(remaining[:boundary].strip())
        remaining = remaining[boundary:].strip()
    return tuple(chunk for chunk in chunks if chunk)


def extract_text_bytes(
    data: bytes,
    name: str,
    *,
    max_segments: int | None = None,
    max_characters: int | None = None,
) -> ExtractedDocument:
    text = data.decode("utf-8-sig").strip()
    if max_characters is not None and len(text) > max_characters:
        raise ExtractionLimitExceeded("Text character budget exceeded.")
    chunks = _chunk_text(text)
    if max_segments is not None and len(chunks) > max_segments:
        raise ExtractionLimitExceeded("Text segment budget exceeded.")
    segments = tuple(
        ExtractedSegment(chunk, None, None, f"text chunk {index}")
        for index, chunk in enumerate(chunks, start=1)
    )
    return ExtractedDocument(name, segments, ())


def extract_document(
    data: bytes,
    name: str,
    *,
    max_pdf_pages: int | None = None,
    sample_pdf_across_document: bool = False,
    max_segments: int | None = None,
    max_characters: int | None = None,
    max_uncompressed_bytes: int | None = None,
    max_archive_members: int | None = None,
) -> ExtractedDocument:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_bytes(
            data,
            name,
            max_pages=max_pdf_pages,
            sample_across_document=sample_pdf_across_document,
            max_segments=max_segments,
            max_characters=max_characters,
            max_uncompressed_bytes=max_uncompressed_bytes,
        )
    if suffix == ".docx":
        return extract_docx_bytes(
            data,
            name,
            max_segments=max_segments,
            max_characters=max_characters,
            max_uncompressed_bytes=max_uncompressed_bytes,
            max_archive_members=max_archive_members,
        )
    if suffix in {".txt", ".md"}:
        return extract_text_bytes(
            data,
            name,
            max_segments=max_segments,
            max_characters=max_characters,
        )
    raise ValueError(f"Unsupported file type: {suffix}")
