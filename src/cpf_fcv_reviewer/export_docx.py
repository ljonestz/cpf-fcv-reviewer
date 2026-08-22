from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .contracts import CurrentEvidenceTier, EvidenceItem, EvidenceLocator, ReviewResult

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x70, 0x85)
BLACK = RGBColor(0, 0, 0)

PAGE_WIDTH_DXA = 9360
LIST_TEXT_INDENT_DXA = 720
LIST_HANGING_DXA = 360
APPLICATION_AUTHOR = "CPF FCV Reviewer"
EVIDENCE_STATUS_LABELS = {
    CurrentEvidenceTier.FULL: "Current evidence established",
    CurrentEvidenceTier.REDUCED: "Current evidence partially established",
    CurrentEvidenceTier.DOCUMENT_LED: "Review based primarily on submitted documents",
}


class EvidenceCompletenessError(ValueError):
    """Raised when a narrative priority area cites unavailable evidence."""


def _set_style_font(style, *, name: str, size: float, color: RGBColor) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.element.rPr.rFonts.set(qn("w:ascii"), name)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), name)


def _set_list_style(style) -> None:
    _set_style_font(style, name="Calibri", size=11, color=BLACK)
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.25)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.167


def _next_numbering_id(numbering, tag: str, attribute: str) -> int:
    values = [int(element.get(qn(attribute))) for element in numbering.findall(qn(tag))]
    return max(values, default=-1) + 1


def _add_numbering_definition(document: Document, *, fmt: str, level_text: str) -> int:
    numbering = document.part.numbering_part.element
    if fmt == "bullet":
        level_text = "•"
    marker = {"decimal": "C0F00001", "bullet": "C0F00002"}[fmt]
    for abstract in numbering.findall(qn("w:abstractNum")):
        nsid = abstract.find(qn("w:nsid"))
        if nsid is None or nsid.get(qn("w:val")) != marker:
            continue
        abstract_id = abstract.get(qn("w:abstractNumId"))
        for number in numbering.findall(qn("w:num")):
            reference = number.find(qn("w:abstractNumId"))
            if reference is not None and reference.get(qn("w:val")) == abstract_id:
                return int(number.get(qn("w:numId")))

    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), marker)
    abstract.append(nsid)
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), level_text)
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(LIST_TEXT_INDENT_DXA))
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(LIST_TEXT_INDENT_DXA))
    indent.set(qn("w:hanging"), str(LIST_HANGING_DXA))
    paragraph_properties.append(indent)
    level.append(paragraph_properties)
    abstract.append(level)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return num_id


def _apply_num_id(paragraph, num_id: int) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    num_properties = paragraph_properties.find(qn("w:numPr"))
    if num_properties is None:
        num_properties = OxmlElement("w:numPr")
        paragraph_properties.append(num_properties)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_properties.extend((ilvl, num_id_element))


def _add_list_paragraph(document: Document, text: str, *, style_name: str, num_id: int):
    paragraph = document.add_paragraph(text, style=style_name)
    _apply_num_id(paragraph, num_id)
    return paragraph


def _add_page_field(paragraph) -> None:
    for field_type in ("begin", "separate", "end"):
        run = paragraph.add_run()
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), field_type)
        run._r.append(field)
        if field_type == "begin":
            instruction_run = paragraph.add_run()
            instruction_run.font.name = "Calibri"
            instruction_run.font.size = Pt(8.5)
            instruction_run.font.color.rgb = MUTED
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = " PAGE "
            instruction_run._r.append(instruction)
        elif field_type == "separate":
            result_run = paragraph.add_run("1")
            result_run.font.name = "Calibri"
            result_run.font.size = Pt(8.5)
            result_run.font.color.rgb = MUTED


def _configure_document(document: Document, *, created_at: datetime) -> tuple[int, int]:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _set_style_font(normal, name="Calibri", size=11, color=BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    timestamp = created_at
    if timestamp.tzinfo is not None:
        timestamp = timestamp.astimezone(UTC).replace(tzinfo=None)
    properties = document.core_properties
    properties.author = APPLICATION_AUTHOR
    properties.last_modified_by = APPLICATION_AUTHOR
    properties.created = timestamp
    properties.modified = timestamp

    title = document.styles["Title"]
    _set_style_font(title, name="Calibri", size=23, color=BLACK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[name]
        _set_style_font(style, name="Calibri", size=size, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True

    for name in ("CPF Decimal List", "CPF Bullet List"):
        if name in document.styles:
            style = document.styles[name]
        else:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_list_style(style)

    revision_num_id = _add_numbering_definition(document, fmt="decimal", level_text="%1.")
    limitation_num_id = _add_numbering_definition(
        document,
        fmt="bullet",
        level_text="•",
    )

    header = section.header.paragraphs[0]
    header.text = "CPF FCV REVIEW | ADVISORY FIRST PASS"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    label = footer.add_run("Volatile-session export | Page ")
    label.font.name = "Calibri"
    label.font.size = Pt(8.5)
    label.font.color.rgb = MUTED
    _add_page_field(footer)

    return revision_num_id, limitation_num_id


def target_text(locator: EvidenceLocator) -> str:
    parts = [locator.document_title]
    if locator.page is not None:
        parts.append(f"page {locator.page}")
    if locator.heading:
        parts.append(locator.heading)
    if locator.element:
        parts.append(locator.element)
    return " | ".join(parts)


def _evidence_type_label(item: EvidenceItem) -> str:
    return {
        "current_context": "Current context",
        "registry_language": "Registry language",
        "user_correction": "User correction",
        "analytical_inference": "Analytical inference",
        "document_fact": "Document source",
    }.get(item.evidence_type, "Source")


def locator_text(item: EvidenceItem) -> str:
    parts = []
    if item.evidence_type == "current_context":
        parts.append(_evidence_type_label(item))
    if item.locator is not None:
        parts.append(target_text(item.locator))
    if item.source_url:
        parts.append(item.source_url)
    if not parts:
        parts.append(_evidence_type_label(item))
    return " | ".join(parts)


def evidence_excerpt(item: EvidenceItem) -> str:
    if item.evidence_type == "current_context":
        return item.text
    return item.locator.excerpt if item.locator is not None else item.text


def validate_evidence_completeness(
    result: ReviewResult,
    evidence: dict[str, EvidenceItem],
) -> None:
    missing = [
        (area.priority_area_id, evidence_id)
        for area in result.priority_areas
        for evidence_id in area.evidence_ids
        if evidence_id not in evidence
    ]
    if missing:
        details = ", ".join(
            f"{priority_area_id}: {evidence_id}"
            for priority_area_id, evidence_id in missing
        )
        raise EvidenceCompletenessError(
            f"Missing evidence for priority area citations: {details}."
        )


def _add_coverage(document: Document, result: ReviewResult) -> None:
    coverage = result.document_coverage
    document.add_paragraph(f"Primary document: {coverage.primary_document}")
    document.add_paragraph(
        "Package documents: " + ", ".join(coverage.package_documents or ("None supplied",))
    )
    document.add_paragraph(
        "Context documents: " + ", ".join(coverage.context_documents or ("None supplied",))
    )
    document.add_paragraph(f"Coverage note: {coverage.coverage_note}")


def _add_reproducibility_metadata(document: Document, result: ReviewResult) -> None:
    metadata = result.metadata
    for label, value in (
        ("Run", metadata.run_id),
        ("Created", metadata.created_at.isoformat()),
        ("Review stage", metadata.review_stage),
        ("Application", metadata.app_release),
        ("Schema", metadata.schema_version),
        ("Rubric", metadata.rubric_version),
        ("Prompts", metadata.prompt_bundle_version),
        ("Model", metadata.model_id),
        ("Diagnostic mode", metadata.diagnostic_mode.value),
        ("Detail level", metadata.detail_level.value),
        ("Current evidence tier", metadata.current_evidence_tier.value),
        (
            "Current evidence limitation",
            metadata.current_evidence_limitation or "None",
        ),
        ("Repair count", str(metadata.repair_count)),
        (
            "Source scan",
            metadata.source_scan_at.isoformat()
            if metadata.source_scan_at is not None
            else "Not recorded",
        ),
        ("Output language", metadata.output_language),
        ("Registry bundle hash", metadata.registry_bundle_hash),
        ("Guidance hash", metadata.guidance_hash),
    ):
        document.add_paragraph(f"{label}: {value}")

    if metadata.parent_run_id is not None:
        document.add_paragraph(f"Parent run: {metadata.parent_run_id}")
    for registry_name, version in sorted(metadata.registry_versions.items()):
        document.add_paragraph(f"Registry {registry_name}: {version}")
    for document_name, fingerprint in sorted(metadata.document_fingerprints.items()):
        document.add_paragraph(f"Document {document_name}: {fingerprint}")
    for prompt_name, prompt_digest in sorted(metadata.prompt_hashes.items()):
        document.add_paragraph(f"Prompt {prompt_name}: {prompt_digest}")
    for outcome in metadata.validation_outcomes:
        document.add_paragraph(f"Validation: {outcome}")
    for correction_id in metadata.correction_ids:
        document.add_paragraph(f"Correction: {correction_id}")


def build_docx(
    result: ReviewResult,
    *,
    evidence: dict[str, EvidenceItem],
    hydrated_referrals: tuple[dict, ...],
) -> bytes:
    validate_evidence_completeness(result, evidence)
    document = Document()
    revision_num_id, limitation_num_id = _configure_document(
        document,
        created_at=result.metadata.created_at,
    )

    document.add_heading("CPF FCV Review", level=0)
    advisory = document.add_paragraph()
    advisory_run = advisory.add_run(
        "Public version. Use public or non-sensitive material only. "
        "AI-assisted advisory first pass. This is not clearance, policy advice, "
        "a compliance finding, an eligibility determination, or an official "
        "classification."
    )
    advisory_run.bold = True
    advisory_run.font.color.rgb = DARK_BLUE

    document.add_heading("Overall assessment", level=1)
    document.add_paragraph(result.overall_read)

    document.add_heading(
        "How the draft responds to the RRA, current FCV dynamics, and the FCV Strategy",
        level=1,
    )
    document.add_paragraph(result.alignment_readout)

    document.add_heading("Priority measures to strengthen the CPF/CEN", level=1)
    if result.revision_summary:
        for item in result.revision_summary:
            _add_list_paragraph(
                document,
                item.title,
                style_name="CPF Decimal List",
                num_id=revision_num_id,
            )
    else:
        document.add_paragraph("No revision summary was returned for this review.")

    document.add_heading("Priority areas for strengthening", level=1)
    if result.priority_areas:
        for area in result.priority_areas:
            document.add_heading(area.heading, level=2)
            document.add_paragraph(area.assessment)
            document.add_paragraph(area.why_it_matters)
            action = document.add_paragraph()
            action.add_run("Recommended action. ").bold = True
            action.add_run(area.recommended_action)
            document.add_paragraph(f"Target: {target_text(area.target_locator)}")
            if area.comment_reference:
                document.add_paragraph(f"Comment addressed: {area.comment_reference}")
            for evidence_id in area.evidence_ids:
                item = evidence.get(evidence_id)
                if item is None:
                    continue
                document.add_paragraph(f"Source: {locator_text(item)}")
                document.add_paragraph(f"Excerpt: {evidence_excerpt(item)}")
    else:
        document.add_paragraph("No priority areas were returned for this review.")

    document.add_paragraph(EVIDENCE_STATUS_LABELS[result.metadata.current_evidence_tier])
    document.add_heading("Limitations and document coverage", level=1)
    if result.limitations:
        for limitation in result.limitations:
            _add_list_paragraph(
                document,
                limitation,
                style_name="CPF Bullet List",
                num_id=limitation_num_id,
            )
    else:
        document.add_paragraph("No additional limitations were recorded.")
    _add_coverage(document, result)

    if hydrated_referrals:
        document.add_heading("Technical appendix", level=1)
        document.add_heading("Institutional referrals", level=2)
        for referral in hydrated_referrals:
            document.add_paragraph(referral["approved_text"])
            document.add_paragraph(f"Registry: {referral['entry_id']} | {referral['version']}")

    document.add_heading("Reproducibility metadata", level=1)
    _add_reproducibility_metadata(document, result)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
