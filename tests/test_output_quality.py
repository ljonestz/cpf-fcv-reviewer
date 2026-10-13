"""Reader-facing regressions found in the Guinea output acceptance review."""
from io import BytesIO
from pathlib import Path

from docx import Document

from cpf_fcv_reviewer.contracts import CoreFinding, CoreQuestion
from cpf_fcv_reviewer.export_docx import build_docx


def test_summary_covers_all_four_questions_without_full_findings(make_valid_result):
    result, evidence = make_valid_result
    result = result.model_copy(update={"core_findings": tuple(
        CoreFinding(question=q, assessment="Detailed reasoning stays in the full note.",
                    summary=f"Brief judgment about {q.value}.", basis="limited")
        for q in CoreQuestion
    )})
    doc = Document(BytesIO(build_docx(
        result, evidence=evidence, hydrated_referrals=(), summary=True,
    )))
    content = "\n".join(p.text for p in doc.paragraphs)
    for q in CoreQuestion:
        assert f"Brief judgment about {q.value}." in content
    assert "Detailed reasoning stays" not in content


def test_export_preserves_arguments_and_renders_emphasis(make_valid_result):
    result, evidence = make_valid_result
    first = ("**The main finding.** Evidence supports it. It matters. "
             "A response follows. Still the same argument.")
    result = result.model_copy(update={"overall_read": first + "\n\nA separate implication."})
    doc = Document(BytesIO(build_docx(result, evidence=evidence, hydrated_referrals=())))
    paragraphs = [p for p in doc.paragraphs if "main finding" in p.text]
    assert len(paragraphs) == 1
    assert paragraphs[0].text.endswith("Still the same argument.")
    assert "**" not in paragraphs[0].text
    assert paragraphs[0].runs[0].bold
    separate = next(p for p in doc.paragraphs if p.text == "A separate implication.")
    assert not any(r.bold for r in separate.runs)


def test_query_input_precedes_output_and_has_no_example_links():
    html = Path("src/cpf_fcv_reviewer/templates/index.html").read_text(encoding="utf-8")
    assert html.index('id="assistant-form"') < html.index('id="assistant-conversation"')
    assert "data-assistant-prompt" not in html
    assert "Try an example" not in html


def test_core_summary_enters_evidence_and_stage_checks(make_valid_result):
    from cpf_fcv_reviewer.validators import result_text, validate_review
    result, evidence = make_valid_result
    result = result.model_copy(update={
        "metadata": result.metadata.model_copy(update={"review_stage": "finalization"}),
        "core_findings": (CoreFinding(
            question=CoreQuestion.IF_THEN_CALIBRATION, basis="limited",
            assessment="The supplied text does not establish conditional responses.",
            summary="Redesign the whole program architecture.",
        ),),
    })
    assert "Redesign the whole program architecture" in result_text(result)
    issues = validate_review(
        result, evidence_ids=set(evidence), evidence=evidence, prohibited_terms=set(),
    )
    assert "stage_overreach" in {issue.code for issue in issues}


def test_if_then_prompt_recognises_existing_adaptation_before_proposals():
    prompt = Path("prompts/review.md").read_text(encoding="utf-8")
    assert 'even when not labelled "If/Then"' in prompt
    assert "Recognise existing" in prompt
    assert 'Suggested scenario' in prompt
    assert 'positive opening' in prompt
    assert 'not existing CPF provisions' in Path("prompts/repair.md").read_text(encoding="utf-8")


def test_browser_preserves_paragraphs_and_escapes_model_markup():
    from test_frontend_download import COMMON_HARNESS, run_node
    completed = run_node(COMMON_HARNESS + r'''
    require(process.argv[1]);
    const hooks = window.__cpfFcvReviewerTestHooks;
    const view = hooks.renderNarrative("**Finding.** Evidence. Implication. More. " +
      "Still one argument.\n\n<script>alert(1)</script>");
    if (view.children.length !== 2) throw Error("Paragraph boundaries changed");
    if (view.textContent.includes("**")) throw Error("Markdown leaked");
    const all = n => [n, ...(n.children || []).flatMap(all)];
    if (all(view).some(n => n.tagName === "script")) throw Error("Unsafe HTML element");
    if (!view.textContent.includes("<script>")) throw Error("Source text lost");
    if (all(view.children[1]).some(n => n.tagName === "strong")) throw Error("Arbitrary emphasis");
    ''')
    assert completed.returncode == 0, completed.stderr


def test_priority_sources_are_present_in_full_export(make_valid_result):
    result, evidence = make_valid_result
    result = result.model_copy(update={"core_findings": (CoreFinding(
        question=CoreQuestion.DRIVERS_PRIORITIES,
        assessment="The driver shapes delivery choices.", basis="limited",
    ),)})
    doc = Document(BytesIO(build_docx(result, evidence=evidence, hydrated_referrals=())))
    paragraphs = [p.text for p in doc.paragraphs]
    start = paragraphs.index("Priority areas for strengthening")
    assert any(p.startswith("Sources:") for p in paragraphs[start:])


def test_new_review_requires_summaries_but_legacy_results_remain_readable(make_valid_result):
    from cpf_fcv_reviewer.validators import validate_review
    result, evidence = make_valid_result
    findings = tuple(CoreFinding(question=q, assessment="Evidence is limited.", basis="limited")
                     for q in CoreQuestion)
    for version, expected in (("2.0.0", False), ("2.1.0", True)):
        candidate = result.model_copy(update={
            "metadata": result.metadata.model_copy(update={"schema_version": version}),
            "core_findings": findings,
        })
        codes = {i.code for i in validate_review(
            candidate, evidence_ids=set(evidence), prohibited_terms=set(),
        )}
        assert ("incomplete_core_findings" in codes) is expected
