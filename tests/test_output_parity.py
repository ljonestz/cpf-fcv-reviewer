from io import BytesIO
from pathlib import Path

from docx import Document

from cpf_fcv_reviewer.app import create_app
from cpf_fcv_reviewer.contracts import CurrentEvidenceTier
from cpf_fcv_reviewer.export_docx import build_docx
from cpf_fcv_reviewer.registry import load_registry_bundle


def test_browser_fields_are_present_in_docx(make_valid_result):
    result, evidence = make_valid_result
    document = Document(BytesIO(build_docx(result, evidence=evidence, hydrated_referrals=())))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert result.overall_read in text
    assert result.alignment_readout in text
    assert result.strategy_readout in text
    for item in result.revision_summary:
        assert item.title not in text
        assert item.priority_area_id not in text
    for area in result.priority_areas:
        for field in (
            area.heading,
            area.assessment,
            area.why_it_matters,
            area.recommended_action,
        ):
            assert field in text
        assert area.target_locator.document_title in text
        assert area.target_locator.heading in text
        assert area.target_locator.element in text
    assert result.limitations[0] not in text
    assert result.document_coverage.coverage_note not in text


def test_canonical_assessments_are_represented_in_docx_with_resolved_evidence(
    make_valid_result,
):
    result, evidence = make_valid_result
    canonical = result.model_dump(mode="json")
    docx_text = "\n".join(
        paragraph.text
        for paragraph in Document(
            BytesIO(build_docx(result, evidence=evidence, hydrated_referrals=()))
        ).paragraphs
    )
    assert canonical["rra_driver_assessments"][0]["driver"] in docx_text
    assert canonical["rra_driver_assessments"][0]["cpf_response"] in docx_text
    assert canonical["rra_driver_assessments"][0]["remaining_gap"] in docx_text
    assert result.strategy_readout in docx_text
    assert canonical["fcv_strategy_assessments"][0]["assessment"] in docx_text
    assert "Anticipate better" in docx_text
    assert "Aligned - High confidence" in docx_text
    assert "Status and confidence" in docx_text
    assert "CPF.docx | Results framework | paragraph 12" in docx_text
    assert "The program will support access." not in docx_text
    assert "rra-1" not in docx_text
    assert "strategy-anticipate-better" not in docx_text
    assert "ev-1" not in docx_text
    assert "Delivery mechanism" not in docx_text
    assert "Result / indicator" not in docx_text
    assert "Gap locus" not in docx_text

def test_docx_and_web_note_share_empty_state_language(make_valid_result):
    result, evidence = make_valid_result
    result = result.model_copy(update={"revision_summary": (), "priority_areas": ()})
    data = build_docx(result, evidence=evidence, hydrated_referrals=())
    text = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)

    assert "No revision summary was returned for this review." not in text
    assert "No priority areas were returned for this review." in text


def test_browser_json_exposes_evidence_status_from_review_metadata(make_valid_result):
    result, evidence = make_valid_result
    registry = load_registry_bundle(
        Path("tests/fixtures/registry_bundle.synthetic.json"),
        allow_synthetic=True,
    )
    app = create_app(
        {"TESTING": True, "START_BACKGROUND_RUNS": False},
        services={"registry_bundle": registry},
    )
    assessment_id = app.extensions["session_store"].create({"status": "created"})
    app.extensions["session_store"].update(
        assessment_id,
        status="complete",
        result=result.model_dump(mode="json"),
        evidence_by_id={key: item.model_dump(mode="json") for key, item in evidence.items()},
    )

    response = app.test_client().get(f"/api/reviews/{assessment_id}/result")

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["metadata"] == result.metadata.model_dump(mode="json")


def test_browser_keeps_evidence_limitations_omitted_from_docx(make_valid_result):
    result, evidence = make_valid_result
    limitation = "Independent current-country research was unavailable."
    result = result.model_copy(
        update={
            "metadata": result.metadata.model_copy(
                update={
                    "current_evidence_tier": CurrentEvidenceTier.DOCUMENT_LED,
                    "current_evidence_limitation": limitation,
                }
            ),
            "limitations": (*result.limitations, limitation),
        }
    )
    registry = load_registry_bundle(
        Path("tests/fixtures/registry_bundle.synthetic.json"),
        allow_synthetic=True,
    )
    app = create_app(
        {"TESTING": True, "START_BACKGROUND_RUNS": False},
        services={"registry_bundle": registry},
    )
    assessment_id = app.extensions["session_store"].create({"status": "created"})
    app.extensions["session_store"].update(
        assessment_id,
        status="complete",
        result=result.model_dump(mode="json"),
        evidence_by_id={key: item.model_dump(mode="json") for key, item in evidence.items()},
    )

    payload = app.test_client().get(f"/api/reviews/{assessment_id}/result").get_json()
    docx_text = "\n".join(
        paragraph.text
        for paragraph in Document(
            BytesIO(build_docx(result, evidence=evidence, hydrated_referrals=()))
        ).paragraphs
    )

    assert payload["metadata"]["current_evidence_tier"] == "document_led"
    assert payload["metadata"]["current_evidence_limitation"] == limitation
    assert "Review based primarily on submitted documents" not in docx_text
    assert "Current evidence tier" not in docx_text
    assert limitation not in docx_text
    assert limitation in payload["limitations"]


def test_detailed_docx_keeps_main_readout_free_of_repeated_evidence_details(
    make_valid_result,
):
    result, evidence = make_valid_result
    document = Document(BytesIO(build_docx(result, evidence=evidence, hydrated_referrals=())))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Evidence and document locations" not in text
    assert "Evidence and reproducibility" not in text
    assert "Source: CPF.docx | Results framework | paragraph 12" not in text
    assert "Excerpt: The program will support access." not in text
    assert "Target: CPF.docx | Results framework | paragraph 12" in text
