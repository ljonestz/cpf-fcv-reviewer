"""Coordinator wording edits are separate from analytical corrections."""
from io import BytesIO

from docx import Document

from cpf_fcv_reviewer.routes import run_assessment


def test_new_intake_uses_coordinator_schema(completed_coordinator_review):
    app, _, assessment_id = completed_coordinator_review
    payload = app.extensions["session_store"].get(assessment_id).payload
    assert payload["result"]["metadata"]["schema_version"] == "2.0.0"
    assert payload["peer_email"]["comments"]


def test_email_overlay_roundtrip_preserves_findings(completed_coordinator_review):
    app, client, assessment_id = completed_coordinator_review
    result = app.extensions["session_store"].get(assessment_id).payload["result"]
    edited = "Dear colleagues,\n\nPlease clarify the delivery logic. <script>literal</script>"
    saved = client.put(f"/api/reviews/{assessment_id}/email", json={
        "text": edited, "base_run_id": result["metadata"]["run_id"],
    })
    assert saved.status_code == 200
    restored = client.get(f"/api/reviews/{assessment_id}/result").json
    assert restored["email_text"] == edited
    assert restored["email_stale"] is False
    assert restored["overall_read"] == result["overall_read"]
    download = client.get(f"/api/reviews/{assessment_id}/export.docx?view=email")
    assert download.status_code == 200
    paragraphs = "\n".join(p.text for p in Document(BytesIO(download.data)).paragraphs)
    assert "<script>literal</script>" in paragraphs


def test_stale_base_cannot_overwrite_email(completed_coordinator_review):
    _, client, assessment_id = completed_coordinator_review
    response = client.put(f"/api/reviews/{assessment_id}/email", json={
        "text": "Draft wording", "base_run_id": "other-run",
    })
    assert response.status_code == 409


def test_substantive_correction_keeps_old_edits_marked_stale(completed_coordinator_review):
    app, client, assessment_id = completed_coordinator_review
    original = client.get(f"/api/reviews/{assessment_id}/result").json
    assert client.put(f"/api/reviews/{assessment_id}/email", json={
        "text": "Coordinator wording to preserve.",
        "base_run_id": original["metadata"]["run_id"],
    }).status_code == 200
    child = client.post(f"/api/reviews/{assessment_id}/corrections", json={
        "text": "The available evidence limits this interpretation.",
    })
    child_id = child.json["assessment_id"]
    run_assessment(app, child_id)
    updated = client.get(f"/api/reviews/{child_id}/result").json
    assert updated["email_text"] == "Coordinator wording to preserve."
    assert updated["email_stale"] is True
    assert updated["peer_email"]["comments"]
    assert updated["metadata"]["parent_run_id"] == assessment_id


def test_correction_can_override_context_framing(completed_coordinator_review):
    app, client, assessment_id = completed_coordinator_review
    response = client.post(f"/api/reviews/{assessment_id}/corrections", json={
        "text": "Use the coordinator's crisis framing.",
        "framing_override": "in_crisis",
    })
    child_id = response.json["assessment_id"]
    run_assessment(app, child_id)
    updated = client.get(f"/api/reviews/{child_id}/result").json
    assert updated["proposed_context_framing"]["category"] == "in_crisis"
    assert updated["proposed_context_framing"]["origin"] == "coordinator_override"


def test_invalid_completed_state_is_rejected_safely(completed_coordinator_review):
    app, client, assessment_id = completed_coordinator_review
    app.extensions["session_store"].update(assessment_id, result={})
    response = client.put(f"/api/reviews/{assessment_id}/email", json={
        "text": "Updated wording.", "base_run_id": assessment_id,
    })
    assert response.status_code == 409


def test_restore_rejects_missing_core_citation(completed_coordinator_review):
    app, client, assessment_id = completed_coordinator_review
    payload = app.extensions["session_store"].get(assessment_id).payload
    result = payload["result"]
    result["core_findings"][0]["evidence_ids"] = ["missing-evidence"]
    app.extensions["session_store"].update(assessment_id, result=result)
    assert client.get(f"/api/reviews/{assessment_id}/result").status_code == 409


def test_email_export_identifies_email_in_subtitle(completed_coordinator_review):
    _, client, assessment_id = completed_coordinator_review
    download = client.get(f"/api/reviews/{assessment_id}/export.docx?view=email")
    text = "\n".join(p.text for p in Document(BytesIO(download.data)).paragraphs)
    assert "PEER-REVIEW EMAIL" in text
    assert "FULL ASSESSMENT" not in text
