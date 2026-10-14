"""The three reader views preserve concise findings, priority order and editing."""
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from test_frontend_download import COMMON_HARNESS, run_node

from cpf_fcv_reviewer.contracts import CoreFinding, CoreQuestion, RevisionSummaryItem
from cpf_fcv_reviewer.export_docx import build_docx

LABELS = (
    "Drivers and priorities", "Critical actions and commitment",
    "Program implications", "Adapting to changing conditions",
)


@pytest.mark.parametrize("save_ok", [False, True])
def test_correction_preserves_unsaved_wording_or_stops(make_valid_result, save_ok):
    result, _ = make_valid_result
    payload = result.model_dump(mode="json") | {"email_text": "Original wording"}
    harness = COMMON_HARNESS + "\nconst complete = " + json.dumps(payload) + ";\n"
    harness += "const saveOk = " + json.dumps(save_ok) + ";\n" + r'''
    (async () => {
      for (const id of ["#peer-email-editor", "#email-status", "#email-warning"])
        nodes[id] = node();
      global.EventSource = class { addEventListener() {} close() {} };
      const writes = [];
      global.fetch = async (url, options) => {
        if (options?.method) writes.push([url, options]);
        return {
          ok: !(url.endsWith("/email") && !saveOk), status: 200,
          json: async () => url.endsWith("/result") ? complete : {
            messages: [], assessment_id: "child", event_url: "events", result_url: "result",
          },
        };
      };
      require(process.argv[1]);
      const hooks = window.__cpfFcvReviewerTestHooks;
      hooks.setAssessmentId("parent");
      await hooks.restoreSavedReview();
      await hooks.loadAssistantHistory();
      nodes["#peer-email-editor"].value = "Unsaved coordinator wording";
      await hooks.updateAssessment("Qualify the finding");
      if (!writes[0]?.[0].endsWith("/email")) throw Error("Wording not saved first");
      if (JSON.parse(writes[0][1].body).text !== "Unsaved coordinator wording")
        throw Error("Wrong wording saved");
      const corrected = writes.some(([url]) => url.endsWith("/corrections"));
      if (corrected !== saveOk) throw Error("Correction proceeded without preserving wording");
      if (!saveOk && nodes["#peer-email-editor"].value !== "Unsaved coordinator wording")
        throw Error("Failed save lost wording");
    })().catch(error => { console.error(error); process.exit(1); });
    '''
    completed = run_node(harness)
    assert completed.returncode == 0, completed.stderr


def test_full_export_has_four_concise_findings(completed_coordinator_review):
    _, client, assessment_id = completed_coordinator_review
    response = client.get(f"/api/reviews/{assessment_id}/export.docx")
    paragraphs = [p.text for p in Document(BytesIO(response.data)).paragraphs]
    paragraphs = paragraphs[:paragraphs.index("Evidence traceability annex")]
    from cpf_fcv_reviewer.review_profiles import CORE_QUESTION_EXPLANATIONS
    for label, explanation in zip(LABELS, CORE_QUESTION_EXPLANATIONS.values(), strict=True):
        assert paragraphs.count(label) == 1
        assert paragraphs[paragraphs.index(label) + 1] == explanation
    assert "RRA driver-to-response assessment" not in paragraphs
    assert any("Where to act" in value for value in paragraphs)


def test_summary_does_not_drop_selected_priorities(make_valid_result):
    result, evidence = make_valid_result
    areas = tuple(
        result.priority_areas[0].model_copy(update={
            "priority_area_id": f"pa-{index}", "heading": f"Priority issue {index}",
        })
        for index in range(1, 6)
    )
    result = result.model_copy(update={
        "priority_areas": areas,
        "revision_summary": tuple(
            RevisionSummaryItem(priority_area_id=a.priority_area_id, title=a.heading)
            for a in areas
        ),
        "core_findings": tuple(
            CoreFinding(
                question=q, assessment="Limited evidence for this question.", basis="limited",
            )
            for q in CoreQuestion
        ),
    })
    doc = Document(BytesIO(build_docx(
        result, evidence=evidence, hydrated_referrals=(), summary=True,
    )))
    main = doc.paragraphs[:next(i for i, p in enumerate(doc.paragraphs)
                                if p.text == "Evidence traceability annex")]
    headings = [p.text for p in main if p.style.name.startswith("Heading")]
    assert [a.heading for a in areas] == [
        text for text in headings if text.startswith("Priority issue")
    ]


def test_intake_and_email_controls_are_optional_and_accessible():
    html = Path("src/cpf_fcv_reviewer/templates/index.html").read_text(encoding="utf-8")
    for element in [
        'id="review-questions"', 'id="previous-comments"', 'id="framing-override"',
        'id="result-tab-email"', 'id="email-panel"', 'id="peer-email-editor"',
        'id="copy-peer-email"', 'id="assistant-input"',
    ]:
        assert element in html
    assert 'aria-controls="email-panel"' in html
    assert 'aria-labelledby="result-tab-email"' in html
    assert 'value="" disabled selected>Select review stage' in html
    for line in html.splitlines():
        if any(f'id="{name}"' in line for name in (
            "review-questions", "previous-comments",
        )):
            assert "required" not in line


@pytest.mark.parametrize("view", ["", "?view=summary"])
def test_notes_include_current_comments_before_grouped_evidence(
    completed_coordinator_review, view,
):
    app, client, assessment_id = completed_coordinator_review
    result = client.get(f"/api/reviews/{assessment_id}/result").json
    original = app.extensions["session_store"].get(assessment_id).payload["result"]
    wording = "Dear colleagues,\n\nCoordinator wording for this review."
    assert client.put(f"/api/reviews/{assessment_id}/email", json={
        "text": wording, "base_run_id": result["metadata"]["run_id"],
    }).status_code == 200
    response = client.get(f"/api/reviews/{assessment_id}/export.docx{view}")
    assert response.status_code == 200
    paragraphs = [p.text for p in Document(BytesIO(response.data)).paragraphs]
    comments = paragraphs.index("Draft review comments")
    annex = paragraphs.index("Evidence traceability annex")
    assert comments < paragraphs.index("Coordinator wording for this review.") < annex
    assert not any(p.startswith(("Sources:", "Target:")) for p in paragraphs[:comments])
    assert any(p.startswith("Sources:") for p in paragraphs[annex:])
    assert app.extensions["session_store"].get(assessment_id).payload["result"] == original


def test_summary_priorities_explain_issue_before_short_response(completed_coordinator_review):
    app, client, assessment_id = completed_coordinator_review
    payload = app.extensions["session_store"].get(assessment_id).payload["result"]
    for area in payload["priority_areas"]:
        area["recommended_action"] += " Verify local feasibility before changing the draft."
    app.extensions["session_store"].update(assessment_id, result=payload)
    result = client.get(f"/api/reviews/{assessment_id}/result").json
    response = client.get(f"/api/reviews/{assessment_id}/export.docx?view=summary")
    paragraphs = [p.text for p in Document(BytesIO(response.data)).paragraphs]
    start = paragraphs.index("Priority measures to strengthen the CPF / CEN")
    end = paragraphs.index("Draft review comments")
    body = "\n".join(paragraphs[start:end])
    for area, item in zip(result["priority_areas"], result["revision_summary"], strict=True):
        assert area["assessment"] in body
        assert area["why_it_matters"] in body
        assert f"Recommended response: {item['title'].rstrip('.')}" in body
        assert area["recommended_action"] not in body
    assert "See the detailed readout for full recommendations" in body


def test_browser_matches_priority_narrative_and_consolidates_traceability(
    completed_coordinator_review,
):
    _, client, assessment_id = completed_coordinator_review
    payload = client.get(f"/api/reviews/{assessment_id}/result").json
    harness = COMMON_HARNESS + "\nconst complete = " + json.dumps(payload) + ";\n"
    harness += r'''
    (async () => {
      nodes["#evidence-annex"] = node();
      global.fetch = async url => ({ok: true, status: 200,
        json: async () => url.endsWith("/result") ? complete : {messages: []}});
      require(process.argv[1]);
      const hooks = window.__cpfFcvReviewerTestHooks;
      hooks.setAssessmentId("saved");
      await hooks.restoreSavedReview();
      const walk = v => [v, ...(v.children || []).flatMap(walk)];
      const cards = walk(nodes["#summary-panel"])
        .filter(v => v.className === "priority-area");
      if (cards.length !== complete.priority_areas.length) throw Error("Lost priorities");
      const first = cards[0];
      const prose = first.children[1].textContent;
      if (!prose.includes(complete.priority_areas[0].assessment)
        || !prose.includes(complete.priority_areas[0].why_it_matters))
        throw Error("Missing explainer");
      const response = first.children[2].textContent;
      if (!response.endsWith(complete.revision_summary[0].title.replace(/[.!?]+$/, "") + "."))
        throw Error("Response not concise");
      if (nodes["#detailed-panel"].textContent.includes("Traceability"))
        throw Error("Inline evidence remains");
      if (!nodes["#evidence-annex"].textContent.includes("Evidence traceability annex"))
        throw Error("Annex missing");
      if (nodes["#evidence-annex"].children[0].tagName !== "details")
        throw Error("Annex not collapsible");
    })().catch(error => {console.error(error); process.exit(1)});
    '''
    completed = run_node(harness)
    assert completed.returncode == 0, completed.stderr
    html = Path("src/cpf_fcv_reviewer/templates/index.html").read_text(encoding="utf8")
    assert html.index('id="evidence-annex"') > html.index('id="email-status"')


@pytest.mark.parametrize("view", ["", "?view=summary"])
@pytest.mark.parametrize("state", ["fallback", "stale"])
def test_note_comments_keep_fallback_or_stale_notice(completed_coordinator_review, view, state):
    app, client, assessment_id = completed_coordinator_review
    store = app.extensions["session_store"]
    if state == "fallback":
        store.update(assessment_id, peer_email=None, email_error="email_draft_unavailable")
        expected = "assembled from the validated recommendations"
    else:
        store.update(assessment_id, email_overlay={
            "text": "Earlier coordinator wording.", "base_run_id": "earlier-run",
        })
        expected = "These wording edits were made for an earlier assessment"
    response = client.get(f"/api/reviews/{assessment_id}/export.docx{view}")
    assert response.status_code == 200
    paragraphs = [p.text for p in Document(BytesIO(response.data)).paragraphs]
    start = paragraphs.index("Draft review comments")
    end = paragraphs.index("Evidence traceability annex")
    assert expected in "\n".join(paragraphs[start:end])
