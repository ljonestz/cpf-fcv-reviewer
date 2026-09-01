from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from cpf_fcv_reviewer import extraction
from cpf_fcv_reviewer.extraction import (
    DiagnosticCoverageUnavailable,
    ExtractionLimitExceeded,
    extract_document,
    extract_docx_bytes,
    extract_pdf_bytes,
    extract_text_bytes,
    require_readable_primary,
    segments_from_pdf_pages,
)


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Strategic context", level=1)
    document.add_paragraph("The CPF identifies localized exclusion.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Outcome"
    table.cell(0, 1).text = "Indicator"
    table.cell(1, 0).text = "Access"
    table.cell(1, 1).text = "People reached"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_pdf(texts: list[str]) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font}
                )
            }
        )
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_docx_segments_retain_heading_and_table_locator():
    extracted = extract_docx_bytes(make_docx(), "CPF.docx")

    paragraph = next(item for item in extracted.segments if "localized" in item.text)
    table = next(item for item in extracted.segments if "People reached" in item.text)

    assert paragraph.heading == "Strategic context"
    assert paragraph.element == "paragraph 2"
    assert table.heading == "Strategic context"
    assert table.element == "table 1 row 2"


def test_detector_docx_segment_budget_stops_element_expansion():
    with pytest.raises(ExtractionLimitExceeded):
        extract_docx_bytes(make_docx(), "CPF.docx", max_segments=2)


def test_diagnostic_coverage_unavailable_is_a_safe_extraction_limit():
    assert issubclass(DiagnosticCoverageUnavailable, ExtractionLimitExceeded)

def test_pdf_segment_budget_stops_page_expansion():
    with pytest.raises(ExtractionLimitExceeded, match="segment"):
        extract_pdf_bytes(
            make_pdf(["First page text", "Second page text"]),
            "CPF.pdf",
            max_segments=1,
        )


def test_pdf_character_budget_stops_page_expansion():
    with pytest.raises(ExtractionLimitExceeded, match="character"):
        extract_pdf_bytes(
            make_pdf(["A page with more extracted text than the budget"]),
            "CPF.pdf",
            max_characters=10,
        )


def test_pdf_uncompressed_stream_budget_is_effective():
    with pytest.raises(ExtractionLimitExceeded, match="PDF stream"):
        extract_pdf_bytes(
            make_pdf(["A decoded page content stream"]),
            "CPF.pdf",
            max_uncompressed_bytes=1,
        )


def test_pdf_segments_use_real_page_numbers_only():
    extracted = segments_from_pdf_pages(
        "CPF.pdf",
        ["First page text", "", "Third page text"],
    )

    assert [segment.page for segment in extracted.segments] == [1, 3]
    assert "page 2 extracted no text" in extracted.warnings


def test_pdf_sampling_warning_uses_exported_suffix(monkeypatch):
    expected_suffix = "conclusions about absence are limited."
    suffix = getattr(extraction, "PDF_SAMPLING_WARNING_SUFFIX", None)
    assert suffix == expected_suffix

    patched_suffix = "sampling warning changed"
    monkeypatch.setattr(extraction, "PDF_SAMPLING_WARNING_SUFFIX", patched_suffix)
    extracted = extract_pdf_bytes(
        make_pdf(["First page", "Second page"]),
        "sample.pdf",
        max_pages=1,
    )

    assert extracted.warnings[-1] == (
        "sample.pdf: sampled 1 of 2 PDF pages; sampling warning changed"
    )


def test_pdf_extraction_can_bound_pages_before_extracting_text(monkeypatch):
    calls = []

    class FakePage:
        def __init__(self, page_number):
            self.page_number = page_number

        def extract_text(self):
            calls.append(self.page_number)
            return f"Page {self.page_number}"

    pages = [FakePage(index) for index in range(1, 6)]
    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda stream: SimpleNamespace(pages=pages),
    )

    extracted = extract_pdf_bytes(b"pdf", "supporting.pdf", max_pages=2)

    assert calls == [1, 2]
    assert [segment.page for segment in extracted.segments] == [1, 2]
    assert extracted.warnings == (
        "supporting.pdf: sampled 2 of 5 PDF pages; "
        "conclusions about absence are limited.",
    )


def test_pdf_sampling_covers_full_page_range_and_warns_about_absence():
    extracted = extract_pdf_bytes(
        make_pdf([f"Page {index}" for index in range(1, 41)]),
        "long-rra.pdf",
        max_pages=12,
        sample_across_document=True,
    )

    expected_pages = (
        1, 2, 3, 4, 5, 11, 16, 22, 28, 33, 39, 40
    )
    assert [segment.page for segment in extracted.segments] == list(expected_pages)
    assert [
        segment.element for segment in extracted.segments
    ] == [f"page {page}" for page in expected_pages]
    assert extracted.warnings[-1] == (
        "long-rra.pdf: sampled 12 of 40 PDF pages; "
        "conclusions about absence are limited."
    )


def test_extracted_text_is_chunked_across_the_full_document():
    text = (
        "COVER PAGE " + "intro " * 900
        + "\n\nIMPLEMENTATION ARRANGEMENTS " + "delivery " * 900
        + "\n\nRESULTS FRAMEWORK " + "indicator " * 900
    )

    extracted = extract_text_bytes(text.encode(), "CPF.txt")

    assert len(extracted.segments) > 1
    assert extracted.segments[0].element == "text chunk 1"
    assert "COVER PAGE" in extracted.segments[0].text
    assert any(
        "RESULTS FRAMEWORK" in segment.text for segment in extracted.segments[1:]
    )
    assert "".join(
        "".join(segment.text for segment in extracted.segments).split()
    ) == "".join(text.split())


def test_extracted_text_chunking_respects_segment_budget():
    text = ("section text " * 1000).encode()

    with pytest.raises(ExtractionLimitExceeded, match="Text segment"):
        extract_text_bytes(text, "CPF.txt", max_segments=1)


def test_empty_primary_is_rejected():
    with pytest.raises(ValueError, match="Primary CPF/CEN is unreadable"):
        require_readable_primary(type("Doc", (), {"segments": (), "warnings": ()})())


@pytest.mark.parametrize("max_pages", (0, -1))
def test_pdf_sampling_with_nonpositive_limit_returns_empty_sample(
    monkeypatch, max_pages
):
    calls = []

    class FakePage:
        def extract_text(self):
            calls.append(True)
            return "Page text"

    pages = [FakePage() for _ in range(5)]
    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda stream: SimpleNamespace(pages=pages),
    )

    extracted = extract_pdf_bytes(
        b"pdf",
        "supporting.pdf",
        max_pages=max_pages,
    )

    assert calls == []
    assert extracted.segments == ()
    assert extracted.warnings == (
        "supporting.pdf: sampled 0 of 5 PDF pages; "
        "conclusions about absence are limited.",
    )


def test_pdf_sampling_retains_empty_sampled_page_warning(monkeypatch):
    class FakePage:
        def __init__(self, page_number):
            self.page_number = page_number

        def extract_text(self):
            if self.page_number == 16:
                return ""
            return f"Page {self.page_number}"

    pages = [FakePage(index) for index in range(1, 41)]
    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda stream: SimpleNamespace(pages=pages),
    )

    extracted = extract_pdf_bytes(
        b"pdf", "long-rra.pdf",
        max_pages=12,
        sample_across_document=True,
    )

    expected_pages = (
        1, 2, 3, 4, 5, 11, 22, 28, 33, 39, 40
    )
    assert [segment.page for segment in extracted.segments] == list(expected_pages)
    assert [
        segment.element for segment in extracted.segments
    ] == [f"page {page}" for page in expected_pages]
    assert "page 16 extracted no text" in extracted.warnings
    assert extracted.warnings[-1] == (
        "long-rra.pdf: sampled 12 of 40 PDF pages; "
        "conclusions about absence are limited."
    )


def test_extract_document_propagates_pdf_sampling_flag():
    extracted = extract_document(
        make_pdf([f"Page {index}" for index in range(1, 6)]),
        "supporting.pdf",
        max_pdf_pages=2,
        sample_pdf_across_document=True,
    )

    assert [segment.page for segment in extracted.segments] == [1, 5]
    assert [segment.element for segment in extracted.segments] == ["page 1", "page 5"]


def test_full_pdf_extraction_preserves_all_extractable_pages_and_blank_warning():
    extracted = extract_document(
        make_pdf([f"Benin RRA page {page}" if page != 50 else "" for page in range(1, 103)]),
        "benin-rra.pdf",
        max_segments=250,
        max_characters=600_000,
        max_uncompressed_bytes=50_000_000,
    )

    assert len(extracted.segments) == 101
    assert extracted.segments[0].page == 1
    assert extracted.segments[-1].page == 102
    assert any(segment.page == 72 for segment in extracted.segments)
    assert "page 50 extracted no text" in extracted.warnings


def test_full_pdf_page_bound_fails_before_sampling_or_text_extraction(monkeypatch):
    calls = []

    class FakePage:
        def extract_text(self):
            calls.append(True)
            return "Page text"

    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda stream: SimpleNamespace(pages=[FakePage() for _ in range(251)]),
    )

    with pytest.raises(ExtractionLimitExceeded, match="page/segment"):
        extract_pdf_bytes(
            b"pdf",
            "oversized-rra.pdf",
            max_segments=250,
        )

    assert calls == []


def test_text_extraction_enforces_uncompressed_byte_bound():
    with pytest.raises(ExtractionLimitExceeded, match="byte budget"):
        extract_text_bytes(
            b"Readable text",
            "oversized-rra.txt",
            max_uncompressed_bytes=1,
        )


def test_pdf_sampling_allows_more_than_segment_bound_for_identification(monkeypatch):
    calls = []

    class FakePage:
        def __init__(self, page_number):
            self.page_number = page_number

        def extract_text(self):
            calls.append(self.page_number)
            return f"Page {self.page_number}"

        def get_contents(self):
            return None

    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda stream: SimpleNamespace(
            pages=[FakePage(page_number) for page_number in range(1, 252)]
        ),
    )

    extracted = extract_pdf_bytes(
        b"pdf",
        "long-rra.pdf",
        max_pages=16,
        sample_across_document=True,
        max_segments=250,
        max_characters=600_000,
        max_uncompressed_bytes=50_000_000,
    )

    assert len(extracted.segments) == 16
    assert calls[0] == 1
    assert calls[-1] == 251
    assert extracted.warnings[-1] == (
        "long-rra.pdf: sampled 16 of 251 PDF pages; "
        "conclusions about absence are limited."
    )
