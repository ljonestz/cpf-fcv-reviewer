import json
from datetime import date
from hashlib import sha256
from io import BytesIO
from pathlib import Path, PurePosixPath
from struct import pack_into, unpack_from
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pydantic import ValidationError

import cpf_fcv_reviewer.extraction as extraction
import cpf_fcv_reviewer.runtime as runtime
from cpf_fcv_reviewer.app import create_app
from cpf_fcv_reviewer.contracts import (
    AssessmentConfidence,
    AssessmentStatus,
    CurrentEvidenceTier,
    DiagnosticEntry,
    DiagnosticMap,
    DiagnosticMode,
    DocumentRole,
    EvidenceLocator,
    EvidencePack,
    FCVStrategicShift,
    FCVStrategyAssessment,
    GapLocus,
    PriorityArea,
    RecommendationScale,
    RRADriverAssessment,
    RevisionSummaryItem,
    SensitivityCategory,
)
from cpf_fcv_reviewer.extraction import (
    DiagnosticCoverageUnavailable,
    DocumentUnreadable,
    ExtractedDocument,
    ExtractedSegment,
    ExtractionLimitExceeded,
)
from cpf_fcv_reviewer.public_research import CurrentContextClaim
from cpf_fcv_reviewer.registry import load_registry_bundle
from cpf_fcv_reviewer.research_controller import ResearchMode, ResearchResult
from cpf_fcv_reviewer.runtime import (
    _select_role_segments,
    _truncate_at_word_boundary,
    build_runtime_services,
)
from cpf_fcv_reviewer.sources import SourceCandidate


def _assessment_mode(payload):
    mode = payload.get("diagnostic_mode")
    if mode is None:
        pack = payload.get("evidence_pack", {})
        metadata = pack.get("metadata", {}) if isinstance(pack, dict) else {}
        mode = metadata.get("diagnostic_mode")
    return getattr(mode, "value", mode)


def test_mechanical_repair_retry_gate_accepts_duplicates_and_rejects_other_issues():
    assert runtime._can_retry_mechanical_repair(
        [
            {"code": "stage_length_overreach"},
            {"code": "stage_length_overreach"},
        ]
    )
    assert not runtime._can_retry_mechanical_repair(
        [{"code": "missing_current_context_support"}]
    )
    assert not runtime._can_retry_mechanical_repair([])


def test_primary_segment_selection_samples_across_long_documents():
    document = ExtractedDocument(
        name="long-cpf.docx",
        segments=tuple(
            ExtractedSegment(
                text=f"Primary segment {index}",
                page=None,
                heading=None,
                element=f"paragraph {index + 1}",
            )
            for index in range(101)
        ),
        warnings=(),
    )

    selected = _select_role_segments(DocumentRole.PRIMARY, (document,), 12)
    selected_text = [segment.text for _, segment, _ in selected]

    assert selected_text[0] == "Primary segment 0"
    assert "Primary segment 55" in selected_text
    assert selected_text[-1] == "Primary segment 100"




@pytest.mark.parametrize(
    ("name", "expected"),
    (
        (
            "oversized-identification.pdf",
            {
                "max_pdf_pages": runtime.OPTIONAL_PDF_SAMPLE_PAGES,
                "sample_pdf_across_document": True,
                "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
                "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
            },
        ),
        (
            "oversized-identification.docx",
            {
                "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
                "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
            },
        ),
        (
            "oversized-identification.txt",
            {
                "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
                "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
            },
        ),
    ),
)
def test_runtime_optional_preflight_applies_diagnostic_bounds_and_excludes_over_limit(
    monkeypatch,
    name,
    expected,
):
    calls = []

    def fake_extract(data, document_name, **kwargs):
        calls.append((document_name, kwargs))
        raise ExtractionLimitExceeded("optional diagnostic bound exceeded")

    monkeypatch.setattr(runtime, "extract_document", fake_extract)
    monkeypatch.setattr(
        runtime, "_has_valid_optional_container", lambda data, suffix: True
    )

    documents, retained_uploads, warnings = runtime._extract_optional_uploads(
        ({"name": name, "bytes": b"oversized"},)
    )

    assert documents == ()
    assert retained_uploads == ()
    assert warnings == (OPTIONAL_UPLOAD_WARNING,)
    assert calls == [(name, expected)]

def test_runtime_optional_pdf_sampling_kwargs_are_pdf_only(monkeypatch):
    calls = []

    def fake_extract(data, name, **kwargs):
        calls.append((name, kwargs))
        return ExtractedDocument(
            name=name,
            segments=(
                ExtractedSegment(
                    text="Usable optional evidence.",
                    page=None,
                    heading=None,
                    element="Paragraph 1",
                ),
            ),
            warnings=(),
        )

    monkeypatch.setattr(runtime, "extract_document", fake_extract)

    documents, retained_uploads, warnings = runtime._extract_optional_uploads(
        (
            {"name": "long-support.pdf", "bytes": b"%PDF-1.7"},
            {"name": "supporting.txt", "bytes": b"plain text"},
        )
    )

    assert [document.name for document in documents] == [
        "long-support.pdf",
        "supporting.txt",
    ]
    assert retained_uploads == (
        (1, {"name": "long-support.pdf", "bytes": b"%PDF-1.7"}),
        (2, {"name": "supporting.txt", "bytes": b"plain text"}),
    )
    assert warnings == ()
    assert calls == [
        (
            "long-support.pdf",
            {
                "max_pdf_pages": 16,
                "sample_pdf_across_document": True,
                "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
                "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
            },
        ),
        (
            "supporting.txt",
            {
                "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
                "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
            },
        ),
    ]


def test_package_segment_budget_is_bounded_and_scales_by_document_count():
    assert runtime._package_segment_budget(()) == 16
    assert runtime._package_segment_budget(tuple(object() for _ in range(3))) == 16
    assert runtime._package_segment_budget(tuple(object() for _ in range(6))) == 18
    assert runtime._package_segment_budget(tuple(object() for _ in range(9))) == 27
    assert runtime._package_segment_budget(tuple(object() for _ in range(20))) == 32


def test_package_selection_gives_nine_documents_three_segments_each():
    documents = tuple(
        ExtractedDocument(
            name=f"package-{index}.txt",
            segments=tuple(
                ExtractedSegment(
                    text=(
                        f"Package {index} segment {segment}"
                        if segment != 1
                        else f"Package {index} results framework"
                    ),
                    page=None,
                    heading=None,
                    element=f"Paragraph {segment + 1}",
                )
                for segment in range(3)
            ),
            warnings=(),
        )
        for index in range(9)
    )

    selected = runtime._select_package_segments(documents)
    counts = {document.name: 0 for document in documents}
    for document, _, _ in selected:
        counts[document.name] += 1

    assert len(selected) == 27
    assert counts == {document.name: 3 for document in documents}


def _assessment_evidence(payload):
    pack = payload.get("evidence_pack")
    if isinstance(pack, dict):
        evidence = pack.get("evidence", ())
        registry_ids = tuple(
            item["evidence_id"]
            for item in evidence
            if item.get("evidence_type") == "registry_language"
            and "PUB-FCV-STRAT-" in item["evidence_id"]
        )
        document_ids = tuple(
            item["evidence_id"]
            for item in evidence
            if item.get("evidence_type") == "document_fact"
        )
    else:
        draft = payload.get("draft", {})
        strategy_rows = draft.get("fcv_strategy_assessments", ())
        registry_ids = tuple(
            evidence_id
            for row in strategy_rows
            for evidence_id in row.get("evidence_ids", ())
        )
        document_ids = tuple(
            evidence_id
            for row in draft.get("rra_driver_assessments", ())
            for evidence_id in row.get("evidence_ids", ())
        )
    return (
        tuple(dict.fromkeys(registry_ids)),
        tuple(dict.fromkeys(document_ids or registry_ids)),
    )


def _test_strategy_assessments(payload):
    registry_ids, _ = _assessment_evidence(payload)
    return tuple(
        FCVStrategyAssessment(
            assessment_id=f"test-strategy-{shift.value}",
            strategic_shift=shift,
            assessment=(
                f"The runtime fixture cannot assess the {shift.value} strategic shift."
            ),
            status=AssessmentStatus.NOT_ASSESSABLE,
            confidence=AssessmentConfidence.LOW,
            gap_locus=None,
            evidence_ids=(registry_ids[index],) if index < len(registry_ids) else (),
        )
        for index, shift in enumerate(FCVStrategicShift)
    )


def _test_rra_assessments(payload):
    if _assessment_mode(payload) != DiagnosticMode.RRA_ALIGNMENT.value:
        return ()
    _, document_ids = _assessment_evidence(payload)
    return (
        RRADriverAssessment(
            assessment_id="test-rra-driver-1",
            driver="The runtime fixture represents a territorial delivery constraint.",
            cpf_response="The CPF fixture includes a bounded response to the constraint.",
            delivery_mechanism="The response uses targeted delivery arrangements.",
            result_or_indicator="The fixture includes a service-access indicator.",
            remaining_gap="Adaptation triggers remain to be specified.",
            status=AssessmentStatus.PARTIALLY_ALIGNED,
            confidence=AssessmentConfidence.MEDIUM,
            gap_locus=GapLocus.MONITORING_ADAPTATION,
            evidence_ids=document_ids[:1],
        ),
    )


def _valid_review_draft(output_type, payload, **values):
    values.setdefault(
        "strategy_readout",
        "The CPF advances prevention and jobs. Operational differentiation remains incomplete.",
    )
    if not values.get("fcv_strategy_assessments"):
        values["fcv_strategy_assessments"] = _test_strategy_assessments(payload)
    if _assessment_mode(payload) == DiagnosticMode.RRA_ALIGNMENT.value:
        if not values.get("rra_driver_assessments"):
            values["rra_driver_assessments"] = _test_rra_assessments(payload)
    else:
        values["rra_driver_assessments"] = ()
    return output_type(**values)


FIXTURE = Path("tests/fixtures/registry_bundle.synthetic.json")
OPTIONAL_UPLOAD_WARNING = (
    "An optional uploaded document could not be read and was excluded."
)


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


RELATIONSHIPS_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
OFFICE_DOCUMENT_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "officeDocument"
)
CUSTOM_XML_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
)
MAIN_DOCUMENT_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)


def _relationship_xml(*relationships: tuple[str, str, str, str | None]) -> bytes:
    body = "".join(
        (
            f'<Relationship Id="{relationship_id}" Type="{relationship_type}" '
            f'Target="{target}"'
            f'{f" TargetMode={target_mode!r}" if target_mode else ""}/>'
        )
        for relationship_id, relationship_type, target, target_mode in relationships
    )
    return (
        f'<Relationships xmlns="{RELATIONSHIPS_NAMESPACE}">{body}</Relationships>'.encode()
    )


def _content_types_xml(
    main_part: str,
    *,
    include_override: bool = True,
    override_part: str | None = None,
    content_type: str = MAIN_DOCUMENT_CONTENT_TYPE,
) -> bytes:
    override = ""
    if include_override:
        override = (
            f'<Override PartName="/{override_part or main_part}" '
            f'ContentType="{content_type}"/>'
        )
    return (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/'
        'vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'{override}<!--X9K7Q--></Types>'
    ).encode()


def _docx_bytes(
    *,
    main_part: str = "custom/main.xml",
    root_relationships: bytes | None = None,
    document_relationships: bytes | None = None,
    include_main_part: bool = True,
    extra_entries: dict[str, bytes] | None = None,
    content_types: bytes | None = None,
) -> bytes:
    if root_relationships is None:
        root_relationships = _relationship_xml(
            ("rId1", OFFICE_DOCUMENT_RELATIONSHIP, main_part, None)
        )
    entries = {
        "[Content_Types].xml": content_types or _content_types_xml(main_part),
        "_rels/.rels": root_relationships,
    }
    if include_main_part:
        entries[main_part] = (
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
            b'2006/main"><w:body><w:p><w:r><w:t>Minimal semantic OOXML evidence.'
            b'</w:t></w:r></w:p>'
            b'</w:body></w:document>'
        )
    if document_relationships is not None:
        main_path = PurePosixPath(main_part)
        relationship_part = main_path.parent / "_rels" / f"{main_path.name}.rels"
        entries[str(relationship_part)] = document_relationships
    entries.update(extra_entries or {})
    return _zip_bytes(entries)


def _mutate_zip_entry_metadata(
    data: bytes,
    entry_name: str,
    *,
    encrypted: bool = False,
    compression_type: int | None = None,
) -> bytes:
    mutated = bytearray(data)
    with ZipFile(BytesIO(data)) as archive:
        info = archive.getinfo(entry_name)
    local_offset = info.header_offset
    assert mutated[local_offset : local_offset + 4] == b"PK\x03\x04"
    if encrypted:
        local_flags = unpack_from("<H", mutated, local_offset + 6)[0]
        pack_into("<H", mutated, local_offset + 6, local_flags | 0x1)
    if compression_type is not None:
        pack_into("<H", mutated, local_offset + 8, compression_type)

    central_offset = 0
    while True:
        central_offset = mutated.find(b"PK\x01\x02", central_offset)
        if central_offset < 0:
            raise AssertionError("ZIP central-directory entry was not found.")
        name_length = unpack_from("<H", mutated, central_offset + 28)[0]
        extra_length = unpack_from("<H", mutated, central_offset + 30)[0]
        comment_length = unpack_from("<H", mutated, central_offset + 32)[0]
        name_start = central_offset + 46
        name_end = name_start + name_length
        if bytes(mutated[name_start:name_end]).decode() == entry_name:
            if encrypted:
                central_flags = unpack_from("<H", mutated, central_offset + 8)[0]
                pack_into("<H", mutated, central_offset + 8, central_flags | 0x1)
            if compression_type is not None:
                pack_into("<H", mutated, central_offset + 10, compression_type)
            return bytes(mutated)
        central_offset = name_end + extra_length + comment_length


def _corrupt_zip_entry(data: bytes, entry_name: str) -> bytes:
    corrupted = bytearray(data)
    with ZipFile(BytesIO(data)) as archive:
        info = archive.getinfo(entry_name)
    name_length = unpack_from("<H", corrupted, info.header_offset + 26)[0]
    extra_length = unpack_from("<H", corrupted, info.header_offset + 28)[0]
    data_start = info.header_offset + 30 + name_length + extra_length
    assert info.compress_size > 2
    corrupted[data_start + info.compress_size // 2] ^= 0xFF
    return bytes(corrupted)


def _python_docx_with_custom_xml_relationship() -> bytes:
    document = Document()
    document.add_paragraph("Python-docx custom XML evidence.")
    buffer = BytesIO()
    document.save(buffer)
    with ZipFile(BytesIO(buffer.getvalue())) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    relationship = (
        f'<Relationship Id="rIdCustomXml" Type="{CUSTOM_XML_RELATIONSHIP}" '
        'Target="../customXml/item1.xml"/>'
    ).encode()
    relationships_name = "word/_rels/document.xml.rels"
    entries[relationships_name] = entries[relationships_name].replace(
        b"</Relationships>",
        relationship + b"</Relationships>",
    )
    entries["customXml/item1.xml"] = b"<context>standard custom XML</context>"
    return _zip_bytes(entries)


VALID_MINIMAL_DOCX = _docx_bytes()
SAFE_PARENT_TARGET_DOCX = _docx_bytes(
    main_part="word/document.xml",
    document_relationships=_relationship_xml(
        ("rId2", CUSTOM_XML_RELATIONSHIP, "../customXml/item1.xml", None)
    ),
    extra_entries={"customXml/item1.xml": b"<context>safe parent target</context>"},
)
PYTHON_DOCX_WITH_CUSTOM_XML = _python_docx_with_custom_xml_relationship()
MALFORMED_OPTIONAL_UPLOADS = (
    ("private-invalid-zip.docx", b"X9K7Q not a zip archive"),
    (
        "private-encrypted-relationships.docx",
        _mutate_zip_entry_metadata(
            VALID_MINIMAL_DOCX,
            "_rels/.rels",
            encrypted=True,
        ),
    ),
    (
        "private-unsupported-compression.docx",
        _mutate_zip_entry_metadata(
            VALID_MINIMAL_DOCX,
            "_rels/.rels",
            compression_type=99,
        ),
    ),
    (
        "private-corrupt-deflate.docx",
        _corrupt_zip_entry(VALID_MINIMAL_DOCX, "_rels/.rels"),
    ),
    (
        "private-missing-content-types.docx",
        _zip_bytes(
            {
                "_rels/.rels": _relationship_xml(
                    ("rId1", OFFICE_DOCUMENT_RELATIONSHIP, "custom/main.xml", None)
                ),
                "custom/main.xml": b"<w:document>X9K7Q</w:document>",
            }
        ),
    ),
    (
        "private-missing-package-relationships.docx",
        _zip_bytes(
            {
                "[Content_Types].xml": (
                    b'<Types xmlns="http://schemas.openxmlformats.org/package/'
                    b'2006/content-types"><!--X9K7Q--></Types>'
                ),
                "custom/main.xml": b"<w:document>X9K7Q</w:document>",
            }
        ),
    ),
    (
        "private-root-without-office-document.docx",
        _docx_bytes(root_relationships=_relationship_xml()),
    ),
    (
        "private-missing-office-document-target.docx",
        _docx_bytes(main_part="custom/X9K7Q-missing.xml", include_main_part=False),
    ),
    (
        "private-malformed-root-relationships.docx",
        _docx_bytes(root_relationships=b"<Relationships>X9K7Q"),
    ),
    (
        "private-missing-main-override.docx",
        _docx_bytes(
            content_types=_content_types_xml(
                "custom/main.xml",
                include_override=False,
            )
        ),
    ),
    (
        "private-wrong-main-content-type.docx",
        _docx_bytes(
            content_types=_content_types_xml(
                "custom/main.xml",
                content_type="application/X9K7Q-wrong",
            )
        ),
    ),
    (
        "private-wrong-main-override-part.docx",
        _docx_bytes(
            content_types=_content_types_xml(
                "custom/main.xml",
                override_part="custom/X9K7Q-wrong.xml",
            )
        ),
    ),
    (
        "private-malformed-content-types.docx",
        _docx_bytes(content_types=b"<Types>X9K7Q"),
    ),
    (
        "private-missing-internal-target.docx",
        _docx_bytes(
            document_relationships=_relationship_xml(
                (
                    "rId2",
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/styles",
                    "X9K7Q-missing.xml",
                    None,
                )
            )
        ),
    ),
    (
        "private-root-escape-target.docx",
        _docx_bytes(
            document_relationships=_relationship_xml(
                (
                    "rId2",
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/styles",
                    "../../X9K7Q.xml",
                    None,
                )
            ),
            extra_entries={"X9K7Q.xml": b"unsafe traversal destination"},
        ),
    ),
    (
        "private-encoded-traversal-target.docx",
        _docx_bytes(
            document_relationships=_relationship_xml(
                ("rId2", CUSTOM_XML_RELATIONSHIP, "%2e%2e/X9K7Q.xml", None)
            ),
            extra_entries={"X9K7Q.xml": b"encoded traversal destination"},
        ),
    ),
    (
        "private-encoded-separator-target.docx",
        _docx_bytes(
            document_relationships=_relationship_xml(
                (
                    "rId2",
                    CUSTOM_XML_RELATIONSHIP,
                    "../customXml%2fX9K7Q.xml",
                    None,
                )
            ),
            extra_entries={"customXml/X9K7Q.xml": b"encoded separator destination"},
        ),
    ),
    ("private-financial-token.pdf", b"X9K7Q-secret-looking-prefix"),
)
MALFORMED_OPTIONAL_IDS = (
    "invalid-zip",
    "encrypted-relationships",
    "unsupported-compression",
    "corrupt-deflate",
    "missing-content-types",
    "missing-package-relationships",
    "root-without-office-document",
    "missing-office-document-target",
    "malformed-root-relationships",
    "missing-main-override",
    "wrong-main-content-type",
    "wrong-main-override-part",
    "malformed-content-types",
    "missing-internal-target",
    "root-escape-target",
    "encoded-traversal-target",
    "encoded-separator-target",
    "invalid-pdf",
)


def test_explicit_services_are_registered_without_runtime_rebuild():
    orchestrator = object()
    bundle = load_registry_bundle(FIXTURE, allow_synthetic=True)

    app = create_app(
        {"TESTING": True, "START_BACKGROUND_RUNS": False},
        services={"review_orchestrator": orchestrator, "registry_bundle": bundle},
    )

    assert app.extensions["review_orchestrator"] is orchestrator
    assert app.extensions["registry_bundle"] is bundle
    assert "session_store" in app.extensions


def production_config(**overrides):
    config = {
        "TESTING": False,
        "ANTHROPIC_API_KEY": "test-key",
        "ANTHROPIC_MODEL_ID": "test-model",
        "REGISTRY_BUNDLE_PATH": str(FIXTURE),
        "REGISTRY_BUNDLE_SHA256": sha256(FIXTURE.read_bytes()).hexdigest(),
        "ALLOW_SYNTHETIC_REGISTRY": False,
        "RESEARCH_MAX_ATTEMPTS": 3,
        "RESEARCH_ATTEMPT_TIMEOUT_SECONDS": 90.0,
        "RESEARCH_TOTAL_BUDGET_SECONDS": 300.0,
        "RESEARCH_MINIMUM_CLAIMS": 4,
        "RESEARCH_MINIMUM_PUBLISHERS": 2,
        "RESEARCH_RETRY_BACKOFF_SECONDS": 1.0,
    }
    config.update(overrides)
    return config


def _current_claims():
    return (
        CurrentContextClaim(
            claim_id="claim-1", text="Exact current finding.", publisher="World Bank",
            source_title="Finding", source_url="https://www.worldbank.org/finding",
            source_date=date(2026, 7, 1), source_type="report", relevance="Relevant.",
            context_kind="structural_dynamic", relationship="establishes",
            licensed_data_required=False,
        ),
        CurrentContextClaim(
            claim_id="claim-2", text="Exact other finding.", publisher="Other source",
            source_title="Other finding", source_url="https://other.example.org/finding",
            source_date=date(2026, 7, 2), source_type="briefing", relevance="Relevant.",
            context_kind="current_development", relationship="corroborates",
            licensed_data_required=False,
        ),
    )


class _InjectedResearchController:
    def __init__(self, result=None):
        self.requests = []
        self.emits = []
        self.allow_document_led = []
        self.result = result

    def run(self, request, emit, *, allow_document_led=False):
        self.requests.append(request)
        self.emits.append(emit)
        self.allow_document_led.append(allow_document_led)
        return self.result or ResearchResult(_current_claims(), {}, 1, CurrentEvidenceTier.FULL)


def test_production_startup_fails_closed_for_missing_registry(tmp_path):
    with pytest.raises(RuntimeError, match="registry bundle"):
        create_app(production_config(REGISTRY_BUNDLE_PATH=str(tmp_path / "missing.json")))


def test_runtime_rejects_registry_hash_mismatch():
    with pytest.raises(RuntimeError, match="hash mismatch"):
        build_runtime_services(
            production_config(
                ALLOW_SYNTHETIC_REGISTRY=True,
                REGISTRY_BUNDLE_SHA256="0" * 64,
            )
        )


def test_production_rejects_synthetic_registry_even_with_valid_hash():
    with pytest.raises(RuntimeError, match="test-only"):
        build_runtime_services(production_config())


def test_production_rejects_registry_without_required_fcv_strategy_entries(tmp_path):
    payload = json.loads(FIXTURE.read_text(encoding="utf-8"))
    payload["synthetic"] = False
    payload["entries"] = [
        entry
        for entry in payload["entries"]
        if entry["entry_id"] != "SYN-PUB-FCV-STRAT-004"
    ]
    path = tmp_path / "production-incomplete.json"
    path.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(RuntimeError, match="required FCV Strategy registry entries"):
        build_runtime_services(
            production_config(
                REGISTRY_BUNDLE_PATH=str(path),
                REGISTRY_BUNDLE_SHA256=sha256(path.read_bytes()).hexdigest(),
            )
        )
def test_runtime_builds_exact_named_step_sequence(monkeypatch):
    class FakeModelGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            self.api_key = api_key
            self.model_id = model_id

    class FakeResearchGateway(FakeModelGateway):
        pass

    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicModelGateway",
        FakeModelGateway,
    )
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeResearchGateway,
    )

    services = build_runtime_services(production_config(ALLOW_SYNTHETIC_REGISTRY=True))

    assert set(services) == {
        "follow_on_gateway",
        "review_orchestrator",
        "registry_bundle",
        "research_controller",
    }
    assert [name for name, _ in services["review_orchestrator"].steps] == [
        "extract",
        "resolve_sources",
        "research",
        "build_evidence",
        "map",
        "review",
        "validate",
        "render",
    ]


def _runtime_services(monkeypatch):
    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    return build_runtime_services(production_config(ALLOW_SYNTHETIC_REGISTRY=True))


def test_runtime_resolve_sources_prefers_direct_original(monkeypatch):
    services = _runtime_services(monkeypatch)
    steps = dict(services["review_orchestrator"].steps)
    original = SourceCandidate("sp-1", "Country RRA.docx", "3", "sharepoint_original", True)
    derived = SourceCandidate("md-1", "Country RRA.md", "3", "derived_copy", True)

    context = steps["resolve_sources"]({"source_candidates": (derived, original)})

    assert context["authoritative_source"] == original


def test_runtime_resolve_sources_uses_upload_fallback(monkeypatch):
    services = _runtime_services(monkeypatch)
    steps = dict(services["review_orchestrator"].steps)
    upload = SourceCandidate("up-1", "Country RRA.docx", "3", "upload", True)

    context = steps["resolve_sources"]({"source_candidates": (upload,)})

    assert context["authoritative_source"] == upload


def test_runtime_cannot_report_completion_without_a_review_result(monkeypatch):
    services = _runtime_services(monkeypatch)
    render = dict(services["review_orchestrator"].steps)["render"]

    with pytest.raises(RuntimeError, match="Review result is unavailable"):
        render({})


def test_runtime_validation_does_not_require_confirmed_priority_response(
    monkeypatch,
    make_valid_result,
):
    result, evidence = make_valid_result
    services = _runtime_services(monkeypatch)
    validate = dict(services["review_orchestrator"].steps)["validate"]
    context = {
        "result": result,
        "evidence_pack": EvidencePack(
            metadata=result.metadata,
            evidence=tuple(evidence.values()),
            diagnostic_entries=(),
        ),
        "payload": {
            "review_focus": "Is the partnership logic credible?",
        },
    }

    validated = validate(context)

    assert "missing_priority_response" not in {
        issue["code"] for issue in validated["validation_issues"]
    }


def test_runtime_validation_passes_incomplete_optional_roles(monkeypatch, make_valid_result):
    captured = {}

    def fake_validate_review(
        result,
        *,
        evidence_ids,
        prohibited_terms,
        incomplete_document_roles=frozenset(),
        registry_entry_ids=None,
    ):
        captured["roles"] = incomplete_document_roles
        captured["registry_entry_ids"] = registry_entry_ids
        return ()

    monkeypatch.setattr(runtime, "validate_review", fake_validate_review)
    reviewed, evidence = make_valid_result
    services = _runtime_services(monkeypatch)
    validate = dict(services["review_orchestrator"].steps)["validate"]
    sampling_warning = (
        "package.pdf: sampled 3 of 9 PDF pages; "
        "conclusions about absence are limited."
    )
    context = {
        "result": reviewed,
        "evidence_pack": EvidencePack(
            metadata=reviewed.metadata,
            evidence=tuple(evidence.values()),
            diagnostic_entries=(),
        ),
        "package_documents": (
            ExtractedDocument("package.pdf", (), (sampling_warning,)),
        ),
        "context_documents": (
            ExtractedDocument("context.txt", (), ("unrelated warning",)),
        ),
    }

    validate(context)

    assert captured["roles"] == frozenset({DocumentRole.PACKAGE})
    assert captured["registry_entry_ids"] == {
        "SYN-REF-001",
        "SYN-PUB-FCV-STRAT-001",
        "SYN-PUB-FCV-STRAT-002",
        "SYN-PUB-FCV-STRAT-003",
        "SYN-PUB-FCV-STRAT-004",
    }


def test_runtime_uses_shared_pdf_sampling_warning_suffix(monkeypatch):
    expected_suffix = "conclusions about absence are limited."
    suffix = getattr(extraction, "PDF_SAMPLING_WARNING_SUFFIX", None)
    assert suffix == expected_suffix
    assert getattr(runtime, "PDF_SAMPLING_WARNING_SUFFIX", None) == suffix

    patched_suffix = "sampling warning changed"
    monkeypatch.setattr(
        runtime, "PDF_SAMPLING_WARNING_SUFFIX", patched_suffix, raising=False
    )
    warning = f"package.pdf: sampled 3 of 9 PDF pages; {patched_suffix}"

    assert runtime._is_incomplete_coverage_warning("package.pdf", warning)


def test_runtime_ignores_lookalike_sampling_warning():
    lookalike_warning = "A package note; conclusions about absence are limited."
    context = {
        "package_documents": (
            ExtractedDocument("package.pdf", (), (lookalike_warning,)),
        ),
        "context_documents": (),
    }

    assert runtime._incomplete_document_roles(context) == frozenset()


def test_runtime_accepts_sampling_warning_for_colon_named_document():
    document_name = "supporting:FY26/report.pdf"
    sampling_warning = (
        f"{document_name}: sampled 3 of 9 PDF pages; "
        "conclusions about absence are limited."
    )
    context = {
        "package_documents": (
            ExtractedDocument(document_name, (), (sampling_warning,)),
        ),
        "context_documents": (),
    }

    assert runtime._incomplete_document_roles(context) == frozenset(
        {DocumentRole.PACKAGE}
    )


def test_runtime_rejects_sampling_warning_for_different_document_name():
    sampling_warning = (
        "other.pdf: sampled 3 of 9 PDF pages; "
        "conclusions about absence are limited."
    )
    context = {
        "package_documents": (
            ExtractedDocument("package.pdf", (), (sampling_warning,)),
        ),
        "context_documents": (),
    }

    assert runtime._incomplete_document_roles(context) == frozenset()


def test_runtime_builds_evidence_and_completes_an_uploaded_review(monkeypatch):
    gateway_calls = []
    current_context_evidence_id = None

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            self.model_id = model_id

        def generate(self, *, prompt_name, payload, output_type):
            nonlocal current_context_evidence_id
            gateway_calls.append(prompt_name)
            if prompt_name == "repair":
                assert {issue["code"] for issue in payload["validation_issues"]} == {
                    "missing_current_context_support",
                    "unknown_institutional_referral",
                }
                available_current_ids = payload["repair_support_evidence_ids"][
                    "current_context"
                ]
                assert current_context_evidence_id in available_current_ids
                assert all(item.startswith("current-") for item in available_current_ids)
                assert payload["repair_support_evidence_ids"]["registry_language"] == [
                    "registry-SYN-PUB-FCV-STRAT-001",
                    "registry-SYN-PUB-FCV-STRAT-002",
                    "registry-SYN-PUB-FCV-STRAT-003",
                    "registry-SYN-PUB-FCV-STRAT-004",
                ]
                repaired_draft = dict(payload["draft"])
                repaired_priority = dict(repaired_draft["priority_areas"][0])
                repaired_priority["evidence_ids"] = (
                    *repaired_priority["evidence_ids"],
                    current_context_evidence_id,
                )
                repaired_draft["priority_areas"] = (repaired_priority,)
                repaired_draft["institutional_referral_ids"] = ()
                return output_type.model_validate(repaired_draft)

            assert prompt_name == "review"
            pack = EvidencePack.model_validate(payload["evidence_pack"])
            evidence_id = next(
                item.evidence_id
                for item in pack.evidence
                if item.evidence_type == "document_fact"
            )
            current_context_evidence_id = next(
                item.evidence_id
                for item in pack.evidence
                if item.evidence_type == "current_context"
            )
            return _valid_review_draft(output_type, payload,
                overall_read="The draft identifies a material delivery constraint.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(
                    RevisionSummaryItem(
                        priority_area_id="area-1",
                        title="Clarify the delivery constraint.",
                    ),
                ),
                priority_areas=(
                    PriorityArea(
                        priority_area_id="area-1",
                        heading="Delivery constraint",
                        assessment="The constraint is described in the uploaded draft.",
                        why_it_matters="It may affect implementation.",
                        recommended_action="Clarify the delivery constraint.",
                        target_locator=EvidenceLocator(
                            document_title="benin-cpf.txt",
                            heading="Paragraph 1",
                            excerpt="Material FCV delivery constraint.",
                        ),
                        recommendation_scale=RecommendationScale.FINE_TUNING,
                        evidence_ids=(evidence_id,),
                        sensitivity=SensitivityCategory.CAUTIOUS,
                        gap_locus=GapLocus.CPF_NARRATIVE,
                    ),
                ),
                institutional_referral_ids=("SYN-REF-999",),
                limitations=("No current RRA was supplied.",),
                coverage_note="The review covers the uploaded CPF.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )
    events = []
    context = services["review_orchestrator"].run(
        {
            "assessment_id": "run-1",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Material FCV delivery constraint. " * 20,
                },
                "package_documents": [],
                "context_documents": [],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: events.append((kind, data)),
    )

    assert context["evidence_pack"].evidence
    assert context["result"].priority_areas[0].evidence_ids == (
        context["evidence_pack"].evidence[0].evidence_id,
        current_context_evidence_id,
    )
    assert context["result"].institutional_referral_ids == ()
    assert gateway_calls == ["review", "repair"]
    assert events[-1][0] == "run_complete"


def test_runtime_excludes_bad_optional_uploads_independently(monkeypatch):
    services = _runtime_services(monkeypatch)
    steps = dict(services["review_orchestrator"].steps)

    def document(name, text):
        return ExtractedDocument(
            name=name,
            segments=(
                ExtractedSegment(
                    text=text,
                    page=None,
                    heading=None,
                    element="Paragraph 1",
                ),
            ) if text else (),
            warnings=(),
        )

    def fake_extract(data, name, **kwargs):
        if name == "benin-cpf.txt":
            return document(name, "Readable primary evidence. " * 10)
        if name == "malformed-package.pdf":
            raise ExtractionLimitExceeded("PRIVATE_FORMAT_DETAIL")
        if name == "empty-context.txt":
            return document(name, "")
        return document(name, "Usable optional evidence.")

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.extract_document", fake_extract)
    payload = {
        "country": "Benin",
        "review_stage": "finalization",
        "detail_level": "standard",
        "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
        "package_documents": [
            {"name": "malformed-package.pdf", "bytes": b"PRIVATE_PACKAGE_TEXT"},
            {"name": "usable-package.txt", "bytes": b"usable package"},
        ],
        "context_documents": [
            {"name": "empty-context.txt", "bytes": b"PRIVATE_CONTEXT_TEXT"},
            {"name": "usable-context.txt", "bytes": b"usable context"},
        ],
        "review_focus": "",
        "corrections": [],
    }

    context = steps["extract"]({"assessment_id": "optional-isolation", "payload": payload})

    assert [item.name for item in context["package_documents"]] == [
        "usable-package.txt"
    ]
    assert [item.name for item in context["context_documents"]] == [
        "usable-context.txt"
    ]
    warning = "An optional uploaded document could not be read and was excluded."
    assert context["extraction_warnings"] == (warning, warning)
    assert not any(
        value in " ".join(context["extraction_warnings"])
        for value in (
            "malformed-package.pdf",
            "empty-context.txt",
            "PRIVATE_FORMAT_DETAIL",
            "PRIVATE_PACKAGE_TEXT",
            "PRIVATE_CONTEXT_TEXT",
        )
    )

    context["research_result"] = ResearchResult(
        _current_claims(), {}, 1, CurrentEvidenceTier.FULL
    )
    context = steps["build_evidence"](context)

    assert set(context["evidence_pack"].metadata.document_fingerprints) == {
        "primary:benin-cpf.txt",
        "package:2:usable-package.txt",
        "context:2:usable-context.txt",
    }


def test_runtime_accepts_minimal_docx_without_document_relationships(monkeypatch):
    with ZipFile(BytesIO(VALID_MINIMAL_DOCX)) as archive:
        assert "custom/_rels/main.xml.rels" not in archive.namelist()

    controller = _InjectedResearchController()
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=object(),
        research_controller=controller,
    )
    steps = dict(services["review_orchestrator"].steps)
    context = steps["extract"](
        {
            "assessment_id": "valid-minimal-docx",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "detail_level": "standard",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Readable primary evidence. " * 10,
                },
                "package_documents": [],
                "context_documents": [
                    {"name": "minimal-context.docx", "bytes": VALID_MINIMAL_DOCX}
                ],
                "review_focus": "",
                "corrections": [],
            },
        }
    )

    assert context["extraction_warnings"] == ()
    assert [document.name for document in context["context_documents"]] == [
        "minimal-context.docx"
    ]
    assert context["context_documents"][0].segments[0].text == (
        "Minimal semantic OOXML evidence."
    )

    context["primary_document"] = ExtractedDocument("benin-cpf.txt", (), ())
    context["_emit"] = lambda *_: None
    context = steps["research"](context)
    assert controller.allow_document_led == [True]
    context = steps["build_evidence"](context)
    assert any(
        item.locator.document_title == "minimal-context.docx"
        and item.text == "Minimal semantic OOXML evidence."
        for item in context["evidence_pack"].evidence
    )


@pytest.mark.parametrize(
    ("name", "docx_bytes", "expected_text"),
    (
        (
            "safe-parent-context.docx",
            SAFE_PARENT_TARGET_DOCX,
            "Minimal semantic OOXML evidence.",
        ),
        (
            "python-docx-context.docx",
            PYTHON_DOCX_WITH_CUSTOM_XML,
            "Python-docx custom XML evidence.",
        ),
    ),
    ids=("explicit-safe-parent", "python-docx-custom-xml"),
)
def test_runtime_accepts_safe_parent_relationship_targets(
    monkeypatch,
    name,
    docx_bytes,
    expected_text,
):
    with ZipFile(BytesIO(docx_bytes)) as archive:
        relationships = archive.read("word/_rels/document.xml.rels")
    assert b'Target="../customXml/item1.xml"' in relationships

    controller = _InjectedResearchController()
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=object(),
        research_controller=controller,
    )
    steps = dict(services["review_orchestrator"].steps)
    context = steps["extract"](
        {
            "payload": {
                "country": "Benin",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Readable primary evidence. " * 10,
                },
                "package_documents": [],
                "context_documents": [{"name": name, "bytes": docx_bytes}],
            }
        }
    )

    assert context["extraction_warnings"] == ()
    assert context["context_documents"][0].segments[0].text == expected_text
    context["primary_document"] = ExtractedDocument("benin-cpf.txt", (), ())
    context["_emit"] = lambda *_: None
    steps["research"](context)
    assert controller.allow_document_led == [True]


@pytest.mark.parametrize(
    ("invalid_name", "invalid_bytes"),
    MALFORMED_OPTIONAL_UPLOADS,
    ids=MALFORMED_OPTIONAL_IDS,
)
def test_runtime_malformed_optional_bytes_are_private_and_valid_context_continues(
    monkeypatch,
    caplog,
    capsys,
    invalid_name,
    invalid_bytes,
):
    secret_marker = "X9K7Q"
    controller = _InjectedResearchController(
        ResearchResult(
            (),
            {},
            1,
            CurrentEvidenceTier.DOCUMENT_LED,
            "Independent current-country research was unavailable.",
        )
    )

    class ModelGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            return _valid_review_draft(output_type, payload,
                overall_read="The uploaded draft can be reviewed with caution.",
                alignment_readout="The draft provides bounded contextual evidence.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the readable submitted documents.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", ModelGateway)
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=controller,
    )
    events = []
    caplog.clear()

    context = services["review_orchestrator"].run(
        {
            "assessment_id": "private-optional-preflight",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Readable primary evidence. " * 10,
                },
                "package_documents": [],
                "context_documents": [
                    {
                        "name": invalid_name,
                        "bytes": invalid_bytes,
                    },
                    {
                        "name": "usable-context.txt",
                        "bytes": b"Usable current-context evidence.",
                    },
                ],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: events.append((kind, data)),
    )

    assert [document.name for document in context["context_documents"]] == [
        "usable-context.txt"
    ]
    assert context["extraction_warnings"] == (OPTIONAL_UPLOAD_WARNING,)
    assert context["evidence_pack"].warnings.count(OPTIONAL_UPLOAD_WARNING) == 1
    assert controller.allow_document_led == [True]
    assert events[-1][0] == "run_complete"
    captured = capsys.readouterr()
    disclosed = " ".join(
        (
            caplog.text,
            captured.out,
            captured.err,
            repr(events),
            context["result"].model_dump_json(),
        )
    )
    assert secret_marker not in disclosed
    assert invalid_name not in disclosed


@pytest.mark.parametrize(
    ("invalid_name", "invalid_bytes"),
    MALFORMED_OPTIONAL_UPLOADS,
    ids=MALFORMED_OPTIONAL_IDS,
)
def test_runtime_malformed_optional_cannot_enable_document_led(
    monkeypatch,
    invalid_name,
    invalid_bytes,
):
    controller = _InjectedResearchController()
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=object(),
        research_controller=controller,
    )
    steps = dict(services["review_orchestrator"].steps)
    primary_without_text = ExtractedDocument("benin-cpf.txt", (), ())

    def allow_document_led_for(context_documents):
        context = steps["extract"](
            {
                "payload": {
                    "country": "Benin",
                    "cpf": {
                        "name": "benin-cpf.txt",
                        "bytes": b"Readable primary evidence. " * 10,
                    },
                    "package_documents": [],
                    "context_documents": context_documents,
                }
            }
        )
        context["primary_document"] = primary_without_text
        context["_emit"] = lambda *_: None
        steps["research"](context)
        return controller.allow_document_led[-1]

    invalid_upload = {
        "name": invalid_name,
        "bytes": invalid_bytes,
    }
    assert allow_document_led_for([invalid_upload]) is False
    assert allow_document_led_for(
        [
            invalid_upload,
            {"name": "usable-context.txt", "bytes": b"Usable context evidence."},
        ]
    ) is True


def test_runtime_optional_extraction_does_not_swallow_programming_defects(monkeypatch):
    services = _runtime_services(monkeypatch)
    extract = dict(services["review_orchestrator"].steps)["extract"]

    def fake_extract(data, name, **kwargs):
        if name == "benin-cpf.txt":
            return ExtractedDocument(
                name=name,
                segments=(
                    ExtractedSegment(
                        text="Readable primary evidence. " * 10,
                        page=None,
                        heading=None,
                        element="Paragraph 1",
                    ),
                ),
                warnings=(),
            )
        raise KeyError("programming defect")

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.extract_document", fake_extract)

    with pytest.raises(KeyError, match="programming defect"):
        extract(
            {
                "payload": {
                    "country": "Benin",
                    "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                    "package_documents": [
                        {"name": "optional.txt", "bytes": b"optional"}
                    ],
                    "context_documents": [],
                }
            }
        )


def test_runtime_primary_document_remains_fail_closed(monkeypatch):
    services = _runtime_services(monkeypatch)
    extract = dict(services["review_orchestrator"].steps)["extract"]
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.extract_document",
        lambda data, name, *, max_pdf_pages=None: ExtractedDocument(name, (), ()),
    )

    with pytest.raises(DocumentUnreadable):
        extract(
            {
                "payload": {
                    "country": "Benin",
                    "cpf": {"name": "benin-cpf.txt", "bytes": b""},
                    "package_documents": [],
                    "context_documents": [],
                }
            }
        )


def test_runtime_preserves_primary_evidence_with_supporting_document(monkeypatch):
    captured = {}

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            pack = EvidencePack.model_validate(payload["evidence_pack"])
            captured["pack"] = pack
            return _valid_review_draft(output_type, payload,
                overall_read="The draft identifies a material delivery constraint.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF and package documents.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )

    context = services["review_orchestrator"].run(
        {
            "assessment_id": "run-with-supporting-document",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Primary CPF delivery constraint. " * 20,
                },
                "package_documents": [
                    {
                        "name": "benin-package.txt",
                        "bytes": b"Supporting package context. " * 10,
                    },
                ],
                "context_documents": [],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    pack = captured["pack"]
    primary_evidence = [
        item
        for item in pack.evidence
        if item.locator and item.locator.document_title == "benin-cpf.txt"
    ]
    package_evidence = [
        item
        for item in pack.evidence
        if item.locator and item.locator.document_title == "benin-package.txt"
    ]
    assert primary_evidence
    assert all(item.document_role == DocumentRole.PRIMARY for item in primary_evidence)
    assert package_evidence
    assert all(item.document_role == DocumentRole.PACKAGE for item in package_evidence)
    assert context["result"].document_coverage.primary_document == "benin-cpf.txt"


def test_runtime_accepts_injected_gateways_and_constructs_configured_research_controller():
    model_gateway = object()
    research_gateway = object()

    services = build_runtime_services(
        production_config(
            ALLOW_SYNTHETIC_REGISTRY=True,
            RESEARCH_MAX_ATTEMPTS=5,
            RESEARCH_TOTAL_BUDGET_SECONDS=42.0,
            RESEARCH_MINIMUM_CLAIMS=6,
            RESEARCH_MINIMUM_PUBLISHERS=3,
            RESEARCH_RETRY_BACKOFF_SECONDS=0.25,
        ),
        model_gateway=model_gateway,
        research_gateway=research_gateway,
    )

    controller = services["research_controller"]
    assert controller.gateway is research_gateway
    assert controller.max_attempts == 5
    assert controller.total_budget_seconds == 42.0
    assert controller.minimum_claims == 6
    assert controller.minimum_publishers == 3
    assert controller._backoff == (0.25,)


def test_default_research_gateway_receives_attempt_timeout(monkeypatch):
    captured = {}

    class FakeModelGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

    class FakeResearchGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds):
            captured["timeout_seconds"] = timeout_seconds

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeModelGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeResearchGateway,
    )

    build_runtime_services(
        production_config(
            ALLOW_SYNTHETIC_REGISTRY=True,
            RESEARCH_ATTEMPT_TIMEOUT_SECONDS=12.5,
        )
    )

    assert captured["timeout_seconds"] == 12.5


def test_runtime_preserves_three_upload_roles_and_focus_in_evidence(monkeypatch):
    captured = {}

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            captured["payload"] = payload
            captured["pack"] = EvidencePack.model_validate(payload["evidence_pack"])
            return _valid_review_draft(output_type, payload,
                overall_read="The draft needs a clearer delivery approach.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                fcv_strategy_assessments=(),
                limitations=(),
                coverage_note="The review covers the CPF, package, and context documents.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )

    services["review_orchestrator"].run(
        {
            "assessment_id": "run-with-three-upload-roles",
            "payload": {
                "country": "Benin",
                "review_stage": "concept_review",
                "detail_level": "in_depth",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Primary CPF delivery constraint. " * 20,
                },
                "package_documents": [
                    {
                        "name": "benin-results.txt",
                        "bytes": b"Package results framework. " * 10,
                    },
                ],
                "context_documents": [
                    {
                        "name": "benin-rra.txt",
                        "bytes": b"Context risk analysis. " * 10,
                    },
                ],
                "review_focus": "Focus on delivery arrangements.",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    pack = captured["pack"]
    by_role = {
        role: [item for item in pack.evidence if item.document_role == role]
        for role in DocumentRole
    }
    assert all(by_role.values())
    assert all(item.evidence_id.startswith("primary-") for item in by_role[DocumentRole.PRIMARY])
    assert all(item.evidence_id.startswith("package-") for item in by_role[DocumentRole.PACKAGE])
    assert all(item.evidence_id.startswith("context-") for item in by_role[DocumentRole.CONTEXT])
    assert set(pack.metadata.document_fingerprints) == {
        "primary:benin-cpf.txt",
        "package:1:benin-results.txt",
        "context:1:benin-rra.txt",
    }
    assert pack.metadata.detail_level.value == "in_depth"
    assert captured["payload"]["review_focus"] == "Focus on delivery arrangements."


def test_runtime_package_deep_section_sampling_covers_later_high_value_sections(monkeypatch):
    captured = {}

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            captured["pack"] = EvidencePack.model_validate(payload["evidence_pack"])
            return _valid_review_draft(output_type, payload,
                overall_read="The draft needs a clearer delivery approach.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                fcv_strategy_assessments=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF and package documents.",
            )

    package_sections = {
        "results.txt": (
            "Results introduction",
            "Results context",
            "Results framework details",
            "Results appendix",
        ),
        "implementation.txt": (
            "Implementation introduction",
            "Implementation context",
            "Implementation arrangements details",
            "Implementation appendix",
        ),
        "monitoring.txt": (
            "Monitoring introduction",
            "Monitoring context",
            "Adaptive management details",
            "Monitoring appendix",
        ),
        "partnerships.txt": (
            "Partnerships introduction",
            "Partnerships context",
            "Partnerships details",
            "Partnerships appendix",
        ),
    }

    def extracted(name, texts):
        return ExtractedDocument(
            name=name,
            segments=tuple(
                ExtractedSegment(
                    text=text,
                    page=None,
                    heading=None,
                    element=f"{name} {index}",
                )
                for index, text in enumerate(texts)
            ),
            warnings=(),
        )

    def fake_extract(data, name, **kwargs):
        if name == "benin-cpf.txt":
            return extracted(
                name,
                tuple(f"Primary segment {index}" for index in range(20)),
            )
        if name in package_sections:
            return extracted(name, package_sections[name])
        return extracted(name, tuple(f"Context segment {index}" for index in range(10)))

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.extract_document", fake_extract)
    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )

    services["review_orchestrator"].run(
        {
            "assessment_id": "run-with-deep-package-sections",
            "payload": {
                "country": "Benin",
                "review_stage": "concept_review",
                "detail_level": "in_depth",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                "package_documents": [
                    {"name": name, "bytes": b"package"}
                    for name in package_sections
                ],
                "context_documents": [],
                "review_focus": "",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    package_items = [
        item
        for item in captured["pack"].evidence
        if item.document_role is DocumentRole.PACKAGE
    ]
    assert {item.locator.document_title for item in package_items} == set(package_sections)
    for title, sections in package_sections.items():
        selected_text = " ".join(
            item.text.casefold()
            for item in package_items
            if item.locator.document_title == title
        )
        assert sections[0].casefold() in selected_text
        assert sections[2].casefold() in selected_text


def test_runtime_package_phase_two_uses_only_additional_markers(monkeypatch):
    captured = {}

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            captured["pack"] = EvidencePack.model_validate(payload["evidence_pack"])
            return _valid_review_draft(output_type, payload,
                overall_read="The draft needs a clearer delivery approach.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                fcv_strategy_assessments=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF and package documents.",
            )

    package_sections = {
        "results.txt": (
            "Results introduction",
            "Generic results background",
            "Results framework section one",
            "Results matrix section two",
            "Results framework section three",
            "Generic results appendix",
        ),
        "implementation.txt": (
            "Implementation introduction",
            "Generic implementation background",
            "Implementation arrangements section one",
            "Delivery arrangements section two",
            "Implementation arrangement section three",
            "Generic implementation appendix",
        ),
        "monitoring.txt": (
            "Monitoring introduction",
            "Generic monitoring background",
            "Adaptive management section one",
            "Risk monitoring section two",
            "Adaptive management section three",
            "Generic monitoring appendix",
        ),
        "partnerships.txt": (
            "Partnerships introduction",
            "Generic stakeholder background",
            "Partnerships section one",
            "Fragility section two",
            "Conflict section three",
            "Generic stakeholder appendix",
        ),
    }

    def extracted(name, texts):
        return ExtractedDocument(
            name=name,
            segments=tuple(
                ExtractedSegment(
                    text=text,
                    page=None,
                    heading=None,
                    element=f"{name} {index}",
                )
                for index, text in enumerate(texts)
            ),
            warnings=(),
        )

    def fake_extract(data, name, **kwargs):
        if name == "benin-cpf.txt":
            return extracted(
                name,
                tuple(f"Primary segment {index}" for index in range(20)),
            )
        if name in package_sections:
            return extracted(name, package_sections[name])
        return extracted(name, tuple(f"Context segment {index}" for index in range(10)))

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.extract_document", fake_extract)
    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )

    services["review_orchestrator"].run(
        {
            "assessment_id": "run-with-phase-two-package-markers",
            "payload": {
                "country": "Benin",
                "review_stage": "concept_review",
                "detail_level": "in_depth",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                "package_documents": [
                    {"name": name, "bytes": b"package"}
                    for name in package_sections
                ],
                "context_documents": [],
                "review_focus": "",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    package_items = [
        item
        for item in captured["pack"].evidence
        if item.document_role is DocumentRole.PACKAGE
    ]
    expected = [
        (title, sections[index])
        for index in (0, 2, 3, 4)
        for title, sections in package_sections.items()
    ]
    assert len(package_items) == 16
    assert [
        (item.locator.document_title, item.text)
        for item in package_items
    ] == expected
    assert all("generic" not in text.casefold() for _, text in expected)


def test_runtime_role_budgets_reserve_context_and_balance_package_documents(monkeypatch):
    captured = {}

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            captured["pack"] = EvidencePack.model_validate(payload["evidence_pack"])
            return _valid_review_draft(output_type, payload,
                overall_read="The draft needs a clearer delivery approach.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                fcv_strategy_assessments=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF and corroborating documents.",
            )

    def extracted(name, role, count):
        return ExtractedDocument(
            name=name,
            segments=tuple(
                ExtractedSegment(
                    text=f"{role} segment {index}",
                    page=None,
                    heading=None,
                    element=f"{role} {index}",
                )
                for index in range(count)
            ),
            warnings=(),
        )

    def fake_extract(data, name, **kwargs):
        if name == "benin-cpf.txt":
            return extracted(name, "primary", 20)
        if name.startswith("package-"):
            return extracted(name, "package", 10)
        return extracted(name, "context", 10)

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.extract_document", fake_extract)
    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )

    services["review_orchestrator"].run(
        {
            "assessment_id": "run-with-saturated-role-budgets",
            "payload": {
                "country": "Benin",
                "review_stage": "concept_review",
                "detail_level": "in_depth",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                "package_documents": [
                    {"name": f"package-{index}.txt", "bytes": b"package"}
                    for index in range(3)
                ],
                "context_documents": [
                    {"name": f"context-{index}.txt", "bytes": b"context"}
                    for index in range(2)
                ],
                "review_focus": "",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    by_role = {
        role: [item for item in captured["pack"].evidence if item.document_role == role]
        for role in DocumentRole
    }
    assert len(by_role[DocumentRole.PRIMARY]) == 12
    assert len(by_role[DocumentRole.PACKAGE]) == 9
    assert len(by_role[DocumentRole.CONTEXT]) == 16
    assert sum(len(items) for items in by_role.values()) == 37
    assert len(
        [item for item in captured["pack"].evidence if item.evidence_type == "current_context"]
    ) == 2
    assert len({item.locator.document_title for item in by_role[DocumentRole.PACKAGE]}) == 3
    assert len({item.locator.document_title for item in by_role[DocumentRole.CONTEXT]}) == 2
    assert len(by_role[DocumentRole.PRIMARY]) > len(by_role[DocumentRole.PACKAGE])


def test_runtime_evidence_truncation_keeps_complete_words(monkeypatch):
    services = _runtime_services(monkeypatch)
    build_evidence = dict(services["review_orchestrator"].steps)["build_evidence"]
    segment_text = (
        ("x " * 299)
        + "obscure "
        + ("y " * 496)
        + "magnified "
        + ("z " * 10)
    )
    context = build_evidence(
        {
            "assessment_id": "word-boundary-evidence",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "detail_level": "standard",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                "package_documents": [],
                "context_documents": [],
                "review_focus": "",
                "corrections": [],
            },
            "primary_document": ExtractedDocument(
                name="benin-cpf.txt",
                segments=(
                    ExtractedSegment(
                        text=segment_text,
                        page=None,
                        heading=None,
                        element="Paragraph 1",
                    ),
                ),
                warnings=(),
            ),
            "research_result": ResearchResult((), {}, 0, CurrentEvidenceTier.FULL),
        }
    )
    evidence = next(
        item
        for item in context["evidence_pack"].evidence
        if item.evidence_type == "document_fact"
    )

    assert len(evidence.text) <= 1600
    assert len(evidence.locator.excerpt) <= 600
    assert evidence.text.endswith("y")
    assert evidence.locator.excerpt.endswith("x")
    assert segment_text[len(evidence.text)].isspace()
    assert segment_text[len(evidence.locator.excerpt)].isspace()


def test_runtime_evidence_truncation_keeps_single_long_token_nonempty():
    truncated = _truncate_at_word_boundary("x" * 1601, 1600)

    assert truncated == "x" * 1600


def test_runtime_bounds_model_visible_corrections_but_preserves_lineage(monkeypatch):
    captured = {}

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            self.model_id = model_id

        def generate(self, *, prompt_name, payload, output_type):
            pack = EvidencePack.model_validate(payload["evidence_pack"])
            captured["pack"] = pack
            return _valid_review_draft(output_type, payload,
                overall_read="The draft requires cautious review.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )
    corrections = [
        {
            "correction_id": f"correction-{index:02d}",
            "created_at": "2026-08-11T18:00:00+00:00",
            "text": "x" * 5000,
            "independently_supported": False,
        }
        for index in range(25)
    ]

    services["review_orchestrator"].run(
        {
            "assessment_id": "run-with-corrections",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"Material FCV delivery constraint. " * 20,
                },
                "package_documents": [],
                "context_documents": [],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": corrections,
            },
        },
        lambda kind, data: None,
    )

    pack = captured["pack"]
    assert len(pack.user_corrections) == 20
    assert all(len(item.text) <= 2000 for item in pack.user_corrections)
    assert len(pack.metadata.correction_ids) == 25


def test_runtime_retries_only_residual_mechanical_repair_issues(monkeypatch):
    repair_payloads = []

    class FakeGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            self.model_id = model_id

        def generate(self, *, prompt_name, payload, output_type):
            if prompt_name == "repair":
                repair_payloads.append(payload)
                overall_read = (
                    "This package is eligible for special treatment."
                    if len(repair_payloads) == 1
                    else "The draft requires cautious review."
                )
            else:
                overall_read = "This package is eligible for special treatment."
            return _valid_review_draft(output_type, payload,
                overall_read=overall_read,
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", FakeGateway)
    monkeypatch.setattr(
        "cpf_fcv_reviewer.runtime.AnthropicPublicResearchGateway",
        FakeGateway,
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=_InjectedResearchController(),
    )

    context = services["review_orchestrator"].run(
        {
            "assessment_id": "run-repair-target",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {
                    "name": "benin-cpf.txt",
                    "bytes": b"SOURCE_SENTINEL material delivery constraint. " * 10,
                },
                "package_documents": [],
                "context_documents": [],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    assert len(repair_payloads) == 2
    assert all(
        payload["forbidden_phrases"] == ("eligible for", "eligible")
        for payload in repair_payloads
    )
    assert all("metadata" not in payload["draft"] for payload in repair_payloads)
    assert all("SOURCE_SENTINEL" not in str(payload) for payload in repair_payloads)
    assert context["result"].metadata.repair_count == 1


def _run_narrow_runtime(monkeypatch, package_text, *, controller=None):
    captured = {}

    class ModelGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            if prompt_name == "diagnostic_map":
                ids = tuple(item["evidence_id"] for item in payload["evidence"])
                return DiagnosticMap(
                    entries=(
                        DiagnosticEntry(
                            entry_id="test-diagnostic-map",
                            short_name="Mapped diagnostic",
                            group="principal_driver",
                            materiality="high",
                            source_evidence_ids=ids,
                            grouping_rationale="The test maps every supplied diagnostic segment.",
                        ),
                    )
                )
            captured["pack"] = EvidencePack.model_validate(payload["evidence_pack"])
            return _valid_review_draft(output_type, payload,
                overall_read="The review identifies a delivery constraint.",
                alignment_readout="The draft partly reflects current context.",
                revision_summary=(), priority_areas=(), institutional_referral_ids=(),
                limitations=(), coverage_note="The review covers the uploaded CPF.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", ModelGateway)
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=controller,
        review_date_provider=lambda: date(2026, 8, 13),
    )
    context = services["review_orchestrator"].run(
        {
            "assessment_id": "narrow-runtime-run",
            "payload": {
                "country": "Benin", "review_stage": "finalization",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"CPF text " * 20},
                "package_documents": [{"name": "package.txt", "bytes": package_text}],
                "context_documents": [], "review_focus": "", "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )
    return context, captured["pack"]


def test_runtime_researches_with_dated_rra_request_and_emitter(monkeypatch):
    controller = _InjectedResearchController()
    context, _ = _run_narrow_runtime(
        monkeypatch,
        b"Benin Risk and Resilience Assessment, March 2025. First segment.",
        controller=controller,
    )

    request = controller.requests[0]
    assert request.mode is ResearchMode.RRA_UPDATE
    assert request.diagnostic_title == "package.txt"
    assert request.diagnostic_date == date(2025, 3, 1)
    assert request.diagnostic_summary == (
        "Benin Risk and Resilience Assessment, March 2025. First segment."
    )
    assert controller.emits[0] is not None
    assert context["uploaded_diagnostic"].name == "package.txt"
    assert context["research_result"].tier is CurrentEvidenceTier.FULL


def test_runtime_uses_holistic_request_without_rra(monkeypatch):
    controller = _InjectedResearchController()
    context, pack = _run_narrow_runtime(
        monkeypatch, b"Package context without a diagnostic marker.", controller=controller
    )

    request = controller.requests[0]
    assert request.mode is ResearchMode.HOLISTIC
    assert request.diagnostic_title is None
    assert request.diagnostic_date is None
    assert context["uploaded_diagnostic"] is None
    assert pack.metadata.diagnostic_mode.value == "limited_framing"


def test_runtime_uses_holistic_request_but_rra_alignment_for_undated_rra(monkeypatch):
    controller = _InjectedResearchController()
    context, pack = _run_narrow_runtime(
        monkeypatch,
        b"Benin Risk and Resilience Assessment. Undated findings.",
        controller=controller,
    )

    request = controller.requests[0]
    assert request.mode is ResearchMode.HOLISTIC
    assert request.diagnostic_title is None
    assert request.diagnostic_date is None
    assert context["uploaded_diagnostic"].publication_date is None
    assert pack.metadata.diagnostic_mode.value == "rra_alignment"


def test_runtime_appends_current_research_exactly_before_review(monkeypatch):
    controller = _InjectedResearchController()
    _, pack = _run_narrow_runtime(
        monkeypatch,
        b"No diagnostic supplied.",
        controller=controller,
    )

    current = [item for item in pack.evidence if item.evidence_type == "current_context"]
    actual = [
        (item.evidence_id, item.text, item.source_url, item.confidence)
        for item in current
    ]
    assert actual == [
        ("current-001", "Exact current finding.", "https://www.worldbank.org/finding", "high"),
        ("current-002", "Exact other finding.", "https://other.example.org/finding", "medium"),
    ]


def test_runtime_appends_registry_evidence_to_model_pack(monkeypatch):
    _, pack = _run_narrow_runtime(
        monkeypatch,
        b"No diagnostic supplied.",
        controller=_InjectedResearchController(),
    )
    bundle = load_registry_bundle(FIXTURE, allow_synthetic=True)

    registry = [item for item in pack.evidence if item.evidence_type == "registry_language"]
    assert [(item.evidence_id, item.text) for item in registry] == [
        (f"registry-{entry.entry_id}", entry.approved_text)
        for entry in bundle.entries
    ]
    assert all(item.confidence == "high" for item in registry)


@pytest.mark.parametrize(
    ("primary_segments", "context_segments", "expected"),
    (
        (("Readable primary",), (), True),
        ((), ("Readable context",), True),
        ((), (), False),
    ),
)
def test_runtime_passes_document_led_only_for_usable_uploaded_evidence(
    monkeypatch,
    primary_segments,
    context_segments,
    expected,
):
    controller = _InjectedResearchController(
        ResearchResult(
            (),
            {},
            1,
            CurrentEvidenceTier.DOCUMENT_LED,
            "Independent current-country research was unavailable.",
        )
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=object(),
        research_controller=controller,
    )
    research = dict(services["review_orchestrator"].steps)["research"]

    def document(name, segments):
        return ExtractedDocument(
            name=name,
            segments=tuple(
                ExtractedSegment(text=text, page=None, heading=None, element="Paragraph 1")
                for text in segments
            ),
            warnings=(),
        )

    context = research(
        {
            "payload": {"country": "Benin"},
            "primary_document": document("benin-cpf.txt", primary_segments),
            "package_documents": (),
            "context_documents": (document("context.txt", context_segments),),
            "_emit": lambda *_: None,
        }
    )

    assert context["research_result"].tier is CurrentEvidenceTier.DOCUMENT_LED
    assert controller.allow_document_led == [expected]


def test_runtime_does_not_turn_uploaded_rra_into_current_context(monkeypatch):
    limitation = "Independent current-country research was unavailable."
    controller = _InjectedResearchController(
        ResearchResult((), {}, 1, CurrentEvidenceTier.DOCUMENT_LED, limitation)
    )
    captured = {}

    class ModelGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            if prompt_name == "diagnostic_map":
                ids = tuple(item["evidence_id"] for item in payload["evidence"])
                return DiagnosticMap(
                    entries=(
                        DiagnosticEntry(
                            entry_id="test-diagnostic-map",
                            short_name="Mapped diagnostic",
                            group="principal_driver",
                            materiality="high",
                            source_evidence_ids=ids,
                            grouping_rationale="The test maps every supplied diagnostic segment.",
                        ),
                    )
                )
            captured["pack"] = EvidencePack.model_validate(payload["evidence_pack"])
            return _valid_review_draft(output_type, payload,
                overall_read="The uploaded draft can be reviewed with caution.",
                alignment_readout="Independent current context was unavailable.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the uploaded CPF and RRA.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", ModelGateway)
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=controller,
    )
    context = services["review_orchestrator"].run(
        {
            "assessment_id": "rra-is-not-current-context",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"CPF text " * 20},
                "package_documents": [
                    {
                        "name": "benin-rra.txt",
                        "bytes": b"Benin Risk and Resilience Assessment. " * 20,
                    }
                ],
                "context_documents": [],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    assert controller.allow_document_led == [True]
    assert not [
        item
        for item in captured["pack"].evidence
        if item.evidence_type == "current_context"
    ]
    assert context["evidence_pack"].metadata.current_evidence_tier is CurrentEvidenceTier.DOCUMENT_LED


def test_runtime_preserves_research_limitation_once_through_repair(monkeypatch):
    authoritative_limitation = (
        "  Independent current-country\n research   was unavailable.  "
    )
    limitation = "Independent current-country research was unavailable."
    unrelated_limitation = "  Model caveat keeps spacing.  "
    controller = _InjectedResearchController(
        ResearchResult(
            (),
            {},
            1,
            CurrentEvidenceTier.DOCUMENT_LED,
            authoritative_limitation,
        )
    )

    class ModelGateway:
        def __init__(self, api_key, model_id, *, timeout_seconds=None):
            pass

        def generate(self, *, prompt_name, payload, output_type):
            if prompt_name == "repair":
                overall_read = "The draft requires cautious review."
                limitations = (
                    "\tIndependent current-country research was   unavailable.\n",
                    unrelated_limitation,
                )
            else:
                overall_read = "This package is eligible for special treatment."
                limitations = (
                    "Independent current-country  research was unavailable.",
                    unrelated_limitation,
                )
            return _valid_review_draft(output_type, payload,
                overall_read=overall_read,
                alignment_readout="Independent current context was unavailable.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=limitations,
                coverage_note="The review covers the uploaded CPF.",
            )

    monkeypatch.setattr("cpf_fcv_reviewer.runtime.AnthropicModelGateway", ModelGateway)
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        research_controller=controller,
    )
    context = services["review_orchestrator"].run(
        {
            "assessment_id": "document-led-repair",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"CPF text " * 20},
                "package_documents": [],
                "context_documents": [],
                "review_focus": "",
                "detail_level": "standard",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    result = context["result"]
    assert context["evidence_pack"].warnings.count(limitation) == 1
    assert result.limitations.count(limitation) == 1
    assert result.limitations == (unrelated_limitation, limitation)
    assert result.metadata.current_evidence_tier is CurrentEvidenceTier.DOCUMENT_LED
    assert result.metadata.current_evidence_limitation == limitation
    assert result.metadata.repair_count == 1


def test_runtime_resolves_duplicate_name_to_context_rra_and_preserves_package_evidence(
    monkeypatch,
):
    calls = []
    captured = {}
    package_sampling_warning = (
        "same.pdf: sampled 8 of 8 PDF pages; "
        "conclusions about absence are limited."
    )
    context_sampling_warning = (
        "same.pdf: sampled 16 of 102 PDF pages; "
        "conclusions about absence are limited."
    )
    primary = ExtractedDocument(
        name="benin-cpf.txt",
        segments=(
            ExtractedSegment(
                "Readable CPF text " * 20, None, None, "paragraph 1"
            ),
        ),
        warnings=(),
    )
    package_bounded = ExtractedDocument(
        name="same.pdf",
        segments=(
            ExtractedSegment(
                "Package-only same-name evidence.", 1, None, "page 1"
            ),
        ),
        warnings=(package_sampling_warning,),
    )
    context_bounded = ExtractedDocument(
        name="same.pdf",
        segments=(
            ExtractedSegment(
                "Benin Risk and Resilience Assessment, March 2025.",
                1,
                None,
                "page 1",
            ),
        ),
        warnings=(context_sampling_warning,),
    )
    context_full = ExtractedDocument(
        name="same.pdf",
        segments=(
            ExtractedSegment("Context RRA page 1.", 1, None, "page 1"),
            ExtractedSegment("Context RRA deep-page issue.", 72, None, "page 72"),
        ),
        warnings=("page 50 extracted no text",),
    )

    def fake_extract(data, name, **kwargs):
        calls.append((data, name, kwargs))
        if name == "benin-cpf.txt":
            return primary
        if kwargs.get("sample_pdf_across_document"):
            return package_bounded if data == b"package-bytes" else context_bounded
        if data != b"context-rra-bytes":
            raise AssertionError("full extraction used unrelated same-name bytes")
        return context_full

    monkeypatch.setattr(runtime, "extract_document", fake_extract)
    monkeypatch.setattr(
        runtime, "_has_valid_optional_container", lambda data, suffix: True
    )

    class FakeGateway:
        def generate(self, *, prompt_name, payload, output_type):
            if prompt_name == "diagnostic_map":
                captured["map_payload"] = payload
                ids = tuple(item["evidence_id"] for item in payload["evidence"])
                return DiagnosticMap(
                    entries=(
                        DiagnosticEntry(
                            entry_id="duplicate-name-map",
                            short_name="Context RRA pages",
                            group="principal_driver",
                            materiality="high",
                            source_evidence_ids=ids,
                            grouping_rationale="Every context RRA page is retained.",
                        ),
                    )
                )
            captured["review_payload"] = payload
            return _valid_review_draft(
                output_type,
                payload,
                overall_read="The review identifies a context delivery constraint.",
                alignment_readout="The mapped context RRA identifies a deep-page issue.",
                revision_summary=(),
                priority_areas=(),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the package and context RRA.",
            )

    controller = _InjectedResearchController()
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=FakeGateway(),
        research_controller=controller,
    )
    context = services["review_orchestrator"].run(
        {
            "assessment_id": "duplicate-name-rra",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "detail_level": "standard",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                "package_documents": [
                    {"name": "same.pdf", "bytes": b"package-bytes"},
                ],
                "context_documents": [
                    {"name": "same.pdf", "bytes": b"context-rra-bytes"},
                ],
                "review_focus": "",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    assert context["uploaded_diagnostic"].source_index == 1
    assert context["diagnostic_document_role"] is DocumentRole.CONTEXT
    assert context["diagnostic_document_position"] == 0
    full_calls = [
        (data, kwargs)
        for data, name, kwargs in calls
        if name == "same.pdf" and not kwargs.get("sample_pdf_across_document")
    ]
    assert full_calls == [
        (
            b"context-rra-bytes",
            {
                "max_pdf_pages": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
                "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
                "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
            },
        ),
    ]
    map_items = captured["map_payload"]["evidence"]
    assert {item["document_role"] for item in map_items} == {"context"}
    assert {item["locator"]["page"] for item in map_items} == {1, 72}
    final_pack = EvidencePack.model_validate(captured["review_payload"]["evidence_pack"])
    package_items = [
        item
        for item in final_pack.evidence
        if item.locator is not None
        and item.locator.document_title == "same.pdf"
        and item.document_role is DocumentRole.PACKAGE
    ]
    assert any("Package-only same-name evidence." in item.text for item in package_items)
    mapped_deep_page = next(
        item
        for item in final_pack.evidence
        if item.evidence_id == "diagnostic-page-072"
    )
    assert mapped_deep_page.document_role is DocumentRole.CONTEXT
    assert mapped_deep_page.locator.page == 72
    assert package_sampling_warning in final_pack.warnings
    assert context_sampling_warning not in final_pack.warnings
    assert "page 50 extracted no text" in final_pack.warnings
    assert controller.requests[0].diagnostic_summary == (
        "Context RRA page 1. Context RRA deep-page issue."
    )


def test_runtime_maps_complete_selected_rra_and_keeps_deep_page_excerpt(monkeypatch):
    calls = []
    captured = {}
    full_segments = tuple(
        ExtractedSegment(
            text=(f"Benin RRA page {page}. " + ("Deep page 102 issue. " if page == 102 else "")),
            page=page,
            heading=None,
            element=f"page {page}",
        )
        for page in range(1, 103)
        if page != 50
    )
    primary = ExtractedDocument(
        name="benin-cpf.txt",
        segments=(ExtractedSegment("Readable CPF text " * 20, None, None, "paragraph 1"),),
        warnings=(),
    )
    bounded = ExtractedDocument(
        name="benin-rra.pdf",
        segments=tuple(
            ExtractedSegment(
                f"Benin Risk and Resilience Assessment page {page}",
                page,
                None,
                f"page {page}",
            )
            for page in range(1, 17)
        ),
        warnings=(
            "benin-rra.pdf: sampled 16 of 102 PDF pages; "
            "conclusions about absence are limited.",
        ),
    )
    full = ExtractedDocument(
        name="benin-rra.pdf",
        segments=full_segments,
        warnings=("page 50 extracted no text",),
    )

    def fake_extract(data, name, **kwargs):
        calls.append((name, kwargs))
        if name == "benin-cpf.txt":
            return primary
        if kwargs.get("sample_pdf_across_document"):
            return bounded
        return full

    monkeypatch.setattr(runtime, "extract_document", fake_extract)

    class FakeGateway:
        def generate(self, *, prompt_name, payload, output_type):
            calls.append((prompt_name, payload))
            if prompt_name == "diagnostic_map":
                captured["map_payload"] = payload
                ids = tuple(item["evidence_id"] for item in payload["evidence"])
                return DiagnosticMap(
                    entries=(
                        DiagnosticEntry(
                            entry_id="diagnostic-all-pages",
                            short_name="Mapped diagnostic pages",
                            group="principal_driver",
                            materiality="high",
                            source_evidence_ids=ids,
                            grouping_rationale="Every supplied extractable page is retained.",
                        ),
                    )
                )
            captured["review_payload"] = payload
            deep_id = "diagnostic-page-102"
            return _valid_review_draft(
                output_type,
                payload,
                overall_read="The CPF needs a clearer response.",
                alignment_readout="The mapped diagnostic identifies a deep-page delivery issue.",
                revision_summary=(),
                priority_areas=(),
                rra_driver_assessments=(
                    RRADriverAssessment(
                        assessment_id="deep-page-driver",
                        driver="The deep-page issue is material.",
                        cpf_response="The CPF partly responds.",
                        delivery_mechanism="Targeted delivery is proposed.",
                        result_or_indicator="A service indicator is included.",
                        remaining_gap="The deep-page issue remains.",
                        status=AssessmentStatus.PARTIALLY_ALIGNED,
                        confidence=AssessmentConfidence.HIGH,
                        gap_locus=GapLocus.MONITORING_ADAPTATION,
                        evidence_ids=(deep_id,),
                    ),
                ),
                institutional_referral_ids=(),
                limitations=(),
                coverage_note="The review covers the complete uploaded RRA.",
            )

    gateway = FakeGateway()
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=gateway,
        research_controller=_InjectedResearchController(),
    )
    context = services["review_orchestrator"].run(
        {
            "assessment_id": "complete-rra-map",
            "payload": {
                "country": "Benin",
                "review_stage": "finalization",
                "detail_level": "standard",
                "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                "package_documents": [{"name": "benin-rra.pdf", "bytes": b"%PDF-1.7"}],
                "context_documents": [],
                "review_focus": "",
                "corrections": [],
            },
        },
        lambda kind, data: None,
    )

    map_payload = captured["map_payload"]
    mapped_items = map_payload["evidence"]
    assert len(mapped_items) == 101
    assert len({item["evidence_id"] for item in mapped_items}) == 101
    assert all("excerpt" not in item["locator"] for item in mapped_items)
    full_call = next(
        kwargs
        for name, kwargs in calls
        if name == "benin-rra.pdf"
        and not kwargs.get("sample_pdf_across_document")
    )
    assert full_call == {
        "max_pdf_pages": runtime.DIAGNOSTIC_MAX_PAGES,
        "max_segments": runtime.DIAGNOSTIC_MAX_PAGES,
        "max_characters": runtime.DIAGNOSTIC_MAX_CHARACTERS,
        "max_uncompressed_bytes": runtime.DIAGNOSTIC_MAX_UNCOMPRESSED_BYTES,
    }
    by_id = {item["evidence_id"]: item for item in mapped_items}
    assert by_id["diagnostic-page-072"]["locator"]["page"] == 72
    assert by_id["diagnostic-page-102"]["locator"]["page"] == 102
    assert by_id["diagnostic-page-102"]["text"].endswith("Deep page 102 issue. ")
    assert "page 50 extracted no text" in context["evidence_pack"].warnings
    coverage_line = (
        "benin-rra.pdf: 102 pages attempted; 101 pages with extractable text; "
        "diagnostic mapping complete."
    )
    assert context["result"].limitations.count(coverage_line) == 1
    assert any(
        "diagnostic mapping complete" in warning
        for warning in context["evidence_pack"].warnings
    )
    final_payload = captured["review_payload"]["evidence_pack"]
    final_ids = {item["evidence_id"] for item in final_payload["evidence"]}
    assert "diagnostic-page-102" in final_ids
    deep_excerpt = next(
        item
        for item in final_payload["evidence"]
        if item["evidence_id"] == "diagnostic-page-102"
    )
    assert len(deep_excerpt["text"]) <= 600
    assert "Deep page 102 issue." in deep_excerpt["text"]
    assert final_payload["diagnostic_entries"]
    assert [name for name, _ in calls if name == "diagnostic_map"] == ["diagnostic_map"]
    assert [name for name, _ in calls if name == "review"] == ["review"]
    assert context["result"].rra_driver_assessments[0].evidence_ids == ("diagnostic-page-102",)


def test_runtime_full_rra_limit_failure_skips_mapping_and_review(monkeypatch):
    calls = []
    primary = ExtractedDocument(
        name="benin-cpf.txt",
        segments=(ExtractedSegment("Readable CPF text " * 20, None, None, "paragraph 1"),),
        warnings=(),
    )
    bounded = ExtractedDocument(
        name="benin-rra.pdf",
        segments=(
            ExtractedSegment("Benin Risk and Resilience Assessment.", 1, None, "page 1"),
        ),
        warnings=(),
    )

    def fake_extract(data, name, **kwargs):
        calls.append((name, kwargs))
        if name == "benin-cpf.txt":
            return primary
        if kwargs.get("sample_pdf_across_document"):
            return bounded
        raise ExtractionLimitExceeded("PDF character budget exceeded.")

    monkeypatch.setattr(runtime, "extract_document", fake_extract)

    class FakeGateway:
        def generate(self, *, prompt_name, payload, output_type):
            calls.append((prompt_name, payload))
            raise AssertionError("No model call should occur after full RRA extraction failure.")

    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=FakeGateway(),
        research_controller=_InjectedResearchController(),
    )

    with pytest.raises(DiagnosticCoverageUnavailable, match="full"):
        services["review_orchestrator"].run(
            {
                "assessment_id": "rra-over-limit",
                "payload": {
                    "country": "Benin",
                    "review_stage": "finalization",
                    "detail_level": "standard",
                    "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                    "package_documents": [
                        {"name": "benin-rra.pdf", "bytes": b"%PDF-1.7"}
                    ],
                    "context_documents": [],
                    "review_focus": "",
                    "corrections": [],
                },
            },
            lambda kind, data: None,
        )

    assert not [
        call
        for call in calls
        if isinstance(call[0], str)
        and call[0] in {"diagnostic_map", "review"}
    ]


def _fresh_diagnostic_map_context(completed_context, final_pack, full_document):
    base_evidence = tuple(
        item
        for item in final_pack.evidence
        if not item.evidence_id.startswith("diagnostic-")
    )
    base_pack = final_pack.model_copy(
        update={"evidence": base_evidence, "diagnostic_entries": ()}
    )
    return {
        "full_diagnostic_document": full_document,
        "diagnostic_document_role": completed_context["diagnostic_document_role"],
        "diagnostic_coverage_warning": completed_context["diagnostic_coverage_warning"],
        "evidence_pack": base_pack,
    }


def _invalid_diagnostic_map_error():
    with pytest.raises(ValidationError) as exc_info:
        DiagnosticMap.model_validate(
            {
                "entries": [
                    {
                        "entry_id": {"raw_secret": "MODEL_OUTPUT_SECRET"},
                        "short_name": "Mapped page",
                        "group": "principal_driver",
                        "materiality": "high",
                        "source_evidence_ids": ["diagnostic-page-001"],
                        "grouping_rationale": "The page is material.",
                        "ignore_previous_instructions": "MODEL_OUTPUT_SECRET",
                    }
                ],
                "unknown_map_field": "MODEL_OUTPUT_SECRET",
            }
        )
    return exc_info.value


class _SequencedDiagnosticMapGateway:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    def generate(self, *, prompt_name, payload, output_type):
        self.calls.append((prompt_name, payload, output_type))
        response = self.responses.pop(0)
        if isinstance(response, BaseException):
            raise response
        if callable(response):
            return response(
                prompt_name=prompt_name,
                payload=payload,
                output_type=output_type,
            )
        return response


def _valid_diagnostic_map(*, payload, **_):
    evidence_ids = tuple(item["evidence_id"] for item in payload["evidence"])
    return DiagnosticMap(
        entries=(
            DiagnosticEntry(
                entry_id="retry-map",
                short_name="Mapped diagnostic pages",
                group="principal_driver",
                materiality="high",
                source_evidence_ids=evidence_ids,
                grouping_rationale="Every supplied page is mapped.",
            ),
        )
    )


def _incomplete_diagnostic_map(*, payload, **_):
    evidence_ids = tuple(item["evidence_id"] for item in payload["evidence"])
    return DiagnosticMap(
        entries=(
            DiagnosticEntry(
                entry_id="incomplete-map",
                short_name="Incomplete diagnostic pages",
                group="principal_driver",
                materiality="high",
                source_evidence_ids=evidence_ids[:-1],
                grouping_rationale="The first supplied pages are mapped.",
            ),
        )
    )


def _coverage_invalid_diagnostic_map(*, payload, **_):
    evidence_ids = tuple(item["evidence_id"] for item in payload["evidence"])
    return DiagnosticMap(
        entries=(
            DiagnosticEntry(
                entry_id="duplicate-entry",
                short_name="Duplicated diagnostic page",
                group="principal_driver",
                materiality="high",
                source_evidence_ids=(
                    evidence_ids[0],
                    evidence_ids[0],
                    "MODEL_OUTPUT_SECRET",
                ),
                grouping_rationale="The model output contains unsafe references.",
            ),
            DiagnosticEntry(
                entry_id="duplicate-entry",
                short_name="Another diagnostic page",
                group="delivery_risk",
                materiality="medium",
                source_evidence_ids=(evidence_ids[0], evidence_ids[1]),
                grouping_rationale="The page is material.",
            ),
        )
    )


def _run_diagnostic_map_step(monkeypatch, gateway, *, package_text=None):
    completed_context, final_pack = _run_narrow_runtime(
        monkeypatch,
        package_text or b"Benin Risk and Resilience Assessment, March 2025.",
        controller=_InjectedResearchController(),
    )
    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=gateway,
        research_controller=_InjectedResearchController(),
    )
    return dict(services["review_orchestrator"].steps)["map"](
        _fresh_diagnostic_map_context(
            completed_context,
            final_pack,
            completed_context["full_diagnostic_document"],
        )
    )


def test_runtime_retries_diagnostic_map_coverage_once_with_expected_diagnostics(
    monkeypatch,
):
    gateway = _SequencedDiagnosticMapGateway(
        _incomplete_diagnostic_map,
        _valid_diagnostic_map,
    )
    package_text = (
        b"Benin Risk and Resilience Assessment, March 2025. "
        + (b"Diagnostic evidence. " * 500)
    )

    context = _run_diagnostic_map_step(
        monkeypatch,
        gateway,
        package_text=package_text,
    )

    assert len(gateway.calls) == 2
    first_call, retry_call = gateway.calls
    assert "coverage_retry" not in first_call[1]
    assert retry_call[1].keys() == first_call[1].keys() | {"coverage_retry"}
    diagnostics = retry_call[1]["coverage_retry"]
    assert diagnostics["missing_material_ids"] == [
        retry_call[1]["material_evidence_ids"][-1]
    ]
    assert diagnostics["duplicated_material_ids"] == []
    assert diagnostics["unknown_model_id_count"] == 0
    assert diagnostics["duplicate_entry_id_count"] == 0
    assert [
        evidence_id
        for entry in context["diagnostic_map"].entries
        for evidence_id in entry.source_evidence_ids
    ] == retry_call[1]["material_evidence_ids"]


def test_runtime_does_not_use_third_diagnostic_map_call_after_schema_then_coverage_failure(
    monkeypatch,
):
    gateway = _SequencedDiagnosticMapGateway(
        _invalid_diagnostic_map_error(),
        _incomplete_diagnostic_map,
        AssertionError("third diagnostic map call is forbidden"),
    )

    with pytest.raises(DiagnosticCoverageUnavailable):
        _run_diagnostic_map_step(monkeypatch, gateway)

    assert len(gateway.calls) == 2
    assert "schema_retry" in gateway.calls[1][1]
    assert "coverage_retry" not in gateway.calls[1][1]


def test_runtime_sanitizes_diagnostic_map_coverage_retry(monkeypatch):
    gateway = _SequencedDiagnosticMapGateway(
        _coverage_invalid_diagnostic_map,
        _valid_diagnostic_map,
    )
    package_text = (
        b"Benin Risk and Resilience Assessment, March 2025. "
        + (b"Diagnostic evidence. " * 500)
    )

    _run_diagnostic_map_step(
        monkeypatch,
        gateway,
        package_text=package_text,
    )

    diagnostics = gateway.calls[1][1]["coverage_retry"]
    assert diagnostics == {
        "missing_material_ids": [gateway.calls[1][1]["material_evidence_ids"][2]],
        "duplicated_material_ids": [gateway.calls[1][1]["material_evidence_ids"][0]],
        "unknown_model_id_count": 1,
        "duplicate_entry_id_count": 1,
        "scaffold": [
            {
                "slot": 1,
                "group": "principal_driver",
                "materiality": "high",
                "source_evidence_ids": [
                    gateway.calls[1][1]["material_evidence_ids"][0]
                ],
            },
            {
                "slot": 2,
                "group": "delivery_risk",
                "materiality": "medium",
                "source_evidence_ids": [
                    gateway.calls[1][1]["material_evidence_ids"][1]
                ],
            },
        ],
    }
    diagnostics_json = json.dumps(diagnostics)
    assert "MODEL_OUTPUT_SECRET" not in diagnostics_json
    assert "unsafe references" not in diagnostics_json
    assert all(
        set(item) == {"slot", "group", "materiality", "source_evidence_ids"}
        for item in diagnostics["scaffold"]
    )


def test_runtime_retries_diagnostic_map_schema_validation_once_with_safe_diagnostics(
    monkeypatch,
):
    gateway = _SequencedDiagnosticMapGateway(
        _invalid_diagnostic_map_error(),
        _valid_diagnostic_map,
    )

    _run_diagnostic_map_step(monkeypatch, gateway)

    assert len(gateway.calls) == 2
    first_call, retry_call = gateway.calls
    assert first_call[0] == retry_call[0] == "diagnostic_map"
    assert first_call[2] is retry_call[2] is DiagnosticMap
    assert "schema_retry" not in first_call[1]
    assert retry_call[1].keys() == first_call[1].keys() | {"schema_retry"}
    diagnostics = retry_call[1]["schema_retry"]
    assert "scaffold" not in diagnostics
    diagnostics_json = json.dumps(diagnostics)
    assert "MODEL_OUTPUT_SECRET" not in diagnostics_json
    assert "ignore_previous_instructions" not in diagnostics_json
    assert "unknown_map_field" not in diagnostics_json
    assert len(diagnostics["issues"]) <= 25
    assert all(len(issue["loc"]) <= 8 for issue in diagnostics["issues"])
    assert all(set(issue) == {"loc", "type"} for issue in diagnostics["issues"])


def test_runtime_fails_closed_after_second_diagnostic_map_schema_validation_error(
    monkeypatch,
):
    first_error = _invalid_diagnostic_map_error()
    second_error = _invalid_diagnostic_map_error()
    gateway = _SequencedDiagnosticMapGateway(first_error, second_error)

    with pytest.raises(DiagnosticCoverageUnavailable) as exc_info:
        _run_diagnostic_map_step(monkeypatch, gateway)

    assert exc_info.value.__cause__ is second_error
    assert len(gateway.calls) == 2


def test_runtime_fails_closed_when_coverage_retry_is_schema_invalid(monkeypatch):
    second_error = _invalid_diagnostic_map_error()
    gateway = _SequencedDiagnosticMapGateway(
        _incomplete_diagnostic_map,
        second_error,
    )

    with pytest.raises(DiagnosticCoverageUnavailable) as exc_info:
        _run_diagnostic_map_step(monkeypatch, gateway)

    assert exc_info.value.__cause__ is second_error
    assert len(gateway.calls) == 2


def test_runtime_fails_closed_when_coverage_retry_remains_incomplete(monkeypatch):
    gateway = _SequencedDiagnosticMapGateway(
        _incomplete_diagnostic_map,
        _incomplete_diagnostic_map,
        AssertionError("third diagnostic map call is forbidden"),
    )

    with pytest.raises(DiagnosticCoverageUnavailable) as exc_info:
        _run_diagnostic_map_step(monkeypatch, gateway)

    assert isinstance(exc_info.value.__cause__, ValueError)
    assert len(gateway.calls) == 2
    assert "coverage_retry" in gateway.calls[1][1]


@pytest.mark.parametrize("error", [RuntimeError("provider failed"), ValueError("bad response")])
def test_runtime_does_not_retry_non_validation_diagnostic_map_errors(monkeypatch, error):
    gateway = _SequencedDiagnosticMapGateway(error, _valid_diagnostic_map)

    with pytest.raises(type(error), match=str(error)):
        _run_diagnostic_map_step(monkeypatch, gateway)

    assert len(gateway.calls) == 1


def test_runtime_serialized_diagnostic_map_budget_is_exact_and_byte_sensitive(monkeypatch):
    completed_context, final_pack = _run_narrow_runtime(
        monkeypatch,
        (
            "Benin Risk and Resilience Assessment, March 2025. "
            + ((chr(0x00e9) + "vidence ") * 100)
        ).encode("utf-8"),
        controller=_InjectedResearchController(),
    )
    full_document = completed_context["full_diagnostic_document"]

    class RecordingMapGateway:
        def __init__(self):
            self.payloads = []

        def generate(self, *, prompt_name, payload, output_type):
            assert prompt_name == "diagnostic_map"
            self.payloads.append(payload)
            ids = tuple(item["evidence_id"] for item in payload["evidence"])
            return DiagnosticMap(
                entries=(
                    DiagnosticEntry(
                        entry_id="budget-map",
                        short_name="Budget test map",
                        group="principal_driver",
                        materiality="high",
                        source_evidence_ids=ids,
                        grouping_rationale="Every supplied page is mapped.",
                    ),
                )
            )

    def map_with(gateway):
        services = build_runtime_services(
            production_config(ALLOW_SYNTHETIC_REGISTRY=True),
            model_gateway=gateway,
            research_controller=_InjectedResearchController(),
        )
        return dict(services["review_orchestrator"].steps)["map"](
            _fresh_diagnostic_map_context(
                completed_context,
                final_pack,
                full_document,
            )
        )

    monkeypatch.setattr(runtime, "DIAGNOSTIC_MAP_MAX_ESTIMATED_INPUT_TOKENS", 160_000)
    initial_gateway = RecordingMapGateway()
    map_with(initial_gateway)
    mapping_payload = initial_gateway.payloads[0]
    serialized = json.dumps(
        mapping_payload,
        ensure_ascii=False,
        separators=(",", ":"),
    )
    character_estimate = (len(serialized) + 2) // 3
    byte_estimate = (len(serialized.encode("utf-8")) + 2) // 3
    assert byte_estimate > character_estimate
    assert all("excerpt" not in item["locator"] for item in mapping_payload["evidence"])

    near_gateway = RecordingMapGateway()
    monkeypatch.setattr(runtime, "DIAGNOSTIC_MAP_MAX_ESTIMATED_INPUT_TOKENS", byte_estimate)
    map_with(near_gateway)
    assert len(near_gateway.payloads) == 1

    failing_gateway = RecordingMapGateway()
    monkeypatch.setattr(
        runtime,
        "DIAGNOSTIC_MAP_MAX_ESTIMATED_INPUT_TOKENS",
        byte_estimate - 1,
    )
    with pytest.raises(DiagnosticCoverageUnavailable, match="safe input budget"):
        map_with(failing_gateway)
    assert failing_gateway.payloads == []



@pytest.mark.parametrize(
    "name",
    (
        "benin-rra.txt",
        "benin-risk-and-resilience-assessment.docx",
        "benin-risk-&-resilience-assessment.txt",
        "accepted-equivalent-diagnostic.pdf",
        "fcv-risk-assessment.txt",
    ),
)
def test_runtime_named_diagnostic_extraction_limit_is_not_silently_excluded(
    monkeypatch,
    name,
):
    def fake_extract(data, document_name, **kwargs):
        raise ExtractionLimitExceeded("named diagnostic bound exceeded")

    monkeypatch.setattr(runtime, "extract_document", fake_extract)
    monkeypatch.setattr(
        runtime, "_has_valid_optional_container", lambda data, suffix: True
    )

    with pytest.raises(DiagnosticCoverageUnavailable, match="named"):
        runtime._extract_optional_uploads(
            ({"name": name, "bytes": b"oversized"},)
        )


def test_runtime_251_page_rra_samples_for_identification_but_fails_full_coverage(
    monkeypatch,
):
    calls = []
    events = []
    primary = ExtractedDocument(
        name="benin-cpf.txt",
        segments=(
            ExtractedSegment("Readable CPF text " * 20, None, None, "paragraph 1"),
        ),
        warnings=(),
    )

    class FakePage:
        def __init__(self, page_number):
            self.page_number = page_number

        def extract_text(self):
            return f"Benin Risk and Resilience Assessment page {self.page_number}."

        def get_contents(self):
            return None

    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda stream: SimpleNamespace(
            pages=[FakePage(page_number) for page_number in range(1, 252)]
        ),
    )
    real_extract = runtime.extract_document

    def fake_extract(data, name, **kwargs):
        calls.append((name, kwargs))
        if name == "benin-cpf.txt":
            return primary
        if kwargs.get("sample_pdf_across_document"):
            return real_extract(data, name, **kwargs)
        raise ExtractionLimitExceeded("PDF page/segment budget exceeded.")

    monkeypatch.setattr(runtime, "extract_document", fake_extract)

    class FakeGateway:
        def generate(self, *, prompt_name, payload, output_type):
            raise AssertionError("No model call should occur after full RRA failure.")

    services = build_runtime_services(
        production_config(ALLOW_SYNTHETIC_REGISTRY=True),
        model_gateway=FakeGateway(),
        research_controller=_InjectedResearchController(),
    )

    with pytest.raises(DiagnosticCoverageUnavailable, match="full"):
        services["review_orchestrator"].run(
            {
                "assessment_id": "rra-251-page-over-limit",
                "payload": {
                    "country": "Benin",
                    "review_stage": "finalization",
                    "detail_level": "standard",
                    "cpf": {"name": "benin-cpf.txt", "bytes": b"primary"},
                    "package_documents": [
                        {"name": "benin-rra-251.pdf", "bytes": b"%PDF-1.7"}
                    ],
                    "context_documents": [],
                    "review_focus": "",
                    "corrections": [],
                },
            },
            lambda kind, data: events.append((kind, data)),
        )

    assert calls[1][0] == "benin-rra-251.pdf"
    assert calls[1][1]["sample_pdf_across_document"] is True
    assert calls[1][1]["max_pdf_pages"] == runtime.OPTIONAL_PDF_SAMPLE_PAGES
    assert calls[2][0] == "benin-rra-251.pdf"
    assert calls[2][1]["max_pdf_pages"] == runtime.DIAGNOSTIC_MAX_PAGES
    assert not [
        event
        for event in events
        if event[0] == "step_complete" and event[1].get("step") == "review"
    ]
    assert events[-1] == (
        "run_failed",
        {"error": "diagnostic_coverage_unavailable"},
    )
